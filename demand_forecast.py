
import argparse
import math
import re
import sys
import json
import calendar
from collections import Counter, defaultdict
from copy import copy
from datetime import datetime, date, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.styles.colors import Color
from openpyxl.styles.cell_style import StyleArray
from openpyxl.worksheet.views import Selection
from openpyxl.utils import get_column_letter


# ======================= EDIT INPUT / OUTPUT PATHS =======================
CONFIG = {
    "inventory": r"E:\Downloads\Automate\nationwise\invenotry.xlsx",
"purchase": r"E:\Downloads\Automate\nationwise\6 month purchase.xlsx",
"sales": r"E:\Downloads\Automate\nationwise\6 month sales.xlsx",
"pricebook": r"E:\Downloads\Automate\nationwise\master price book.xlsx",
"purchase_append": r"E:\Downloads\Automate\nationwise\crrnt purchase warehouse.xlsx",
"sales_append": r"E:\Downloads\Automate\nationwise\crrnt sales warehouse.xlsx",
"stores_sales": r"E:\Downloads\Automate\nationwise\6 months pdi sales.xlsx",
"open_po": r"E:\Downloads\Automate\nationwise\open PO's.xlsx",
"planning_inputs": "",
"history": "",
"output": r"E:\Downloads\Automate\Warehouse_Monthly_Sales_new_J.xlsx",
"output_ordering": r"E:\Downloads\Automate\Warehouse_Ordering_Sheet_J.xlsx"
}

STORES_BASE_YEAR = 2026  # Year for month-only PDI rows (change for a new base year)
STORES_APPEND_YEAR = 2026  # Year for month-only append rows
STORES_APPEND_COVERAGE = ("", "")  # Dated daily exports: complete inclusive YYYY-MM-DD range
AS_OF_DATE = ""  # YYYY-MM-DD; blank = today's date in Pakistan
# Optional declared sales coverage if QB report labels cannot be read.
# Example: ("2026-04-01", "2026-09-30"). Blank uses report/date evidence.
SALES_BASE_COVERAGE = ("", "")
# Optional explicit coverage for append exports, inclusive YYYY-MM-DD.
# Normally the QB report period is detected automatically. For partial reports,
# do NOT use date replacement: provide a complete daily report instead.
APPEND_COVERAGE = {"purchase": ("", ""), "sales": ("", "")}
# None selects the first worksheet containing the expected QB headers.
SHEET_NAMES = {"inventory": None, "purchase": None, "sales": None}
PRICEBOOK_SHEET = "BTCM"
SORT_CLEAN_SHEETS = True
# False = do NOT write the copies of the original exports ("... raw", "current ...") into the output.
# They are only used while the script runs; skipping them makes saving much faster and the file far smaller
# (the 245k-row sales raw tab alone is ~2.4 million cells). Your original files are untouched either way.
INCLUDE_RAW_TABS = False
# ---------------- ORDERING SHEET (tab 2 of the second workbook) ----------------
# Add, delete or reorder columns here. Every name must be a column header of 'final sheet'
# (Department, Brand, box upc, master upc, description, Pack of, QTY, Negative QTY flag, PB description,
#  WH last 30 days, Daily demand, Days cover, Stock at arrival, Incoming QTY, Suggested order QTY, Priority,
#  ABC class, Item class, Planning notes ... any header of the final sheet).
# Use ("final sheet header", "Heading to show") to rename a column in the ordering sheet.
# Two extra columns are built into the ordering sheet and are not taken from the final sheet:
#   "Adjustment"      = blank green input cell: type the COMPLETE quantity you want to order (e.g. suggested 10,
#                       you want 15 -> type 15). Leave it blank to keep the suggestion; type 0 to order nothing.
#   "Final order QTY" = formula: the Adjustment if you typed one, otherwise the Suggested order QTY
ORDERING_COLUMNS = ["Department", "Brand", "box upc", "master upc", "description", "PB description",
                    "Pack of", "QTY", "Incoming QTY", "Suggested order QTY", "Adjustment", "Final order QTY", "Priority",
                    "Last buying price", "Last vendor", "Lowest buying price", "Lowest price vendor"]
ORDERING_BORDERS = True            # thin borders around every cell of the ordering sheet (header and data)
ORDERING_ONLY_ORDER_ROWS = True    # True = the ordering sheet lists only the SKUs whose Priority is in ORDERING_PRIORITIES
ORDERING_PRIORITIES = ["Urgent - order", "Order"]   # add "Urgent - expedite incoming" etc. if you want them listed too
ORDERING_SHEET_NAME = "ordering sheet"
FINAL_SALES_COLUMN_WIDTH = 18  # Fixed width for every WH/stores month column
# Fixed leading columns of "final sheet": Status, Department, Brand, box upc, master upc,
# description, Pack of, QTY, PB description. Month columns start right after them.
OPEN_PO_STATUS = "Active - open PO (was inactive)"
FINAL_BASE_HEADERS = ["Status", "Department", "Brand", "box upc", "master upc", "description", "Pack of", "QTY", "Negative QTY flag", "PB description"]
BASE_COLS = len(FINAL_BASE_HEADERS)
DESC_COL = FINAL_BASE_HEADERS.index("description") + 1
QTY_COL_LETTER = "H"  # QTY column of FINAL_BASE_HEADERS
# Additional verified suffixes can be added here, e.g. "OLD UPC".
# Unknown trailing words are NOT automatically removed: they may be flavor/model.
EXTRA_DESCRIPTION_SUFFIXES = []

# ======================= DEMAND PLANNING SETTINGS =======================
# These become the DEFAULT values written into the editable green input
# cells of every row in 'final sheet'. You can still change any cell in
# Excel afterwards and the formulas recalculate.
DEMAND = {
    "lead_time": 3,               # vendor delivery days
    "reorder_below_days": 10,     # order when days of cover is BELOW this (base value, see class override)
    "target_days": 20,            # stock wanted AFTER the order arrives
    "window_days": 30,            # length of each of the 3 rolling sales windows
    "weights": (0.5, 0.3, 0.2),   # newest window first; must add up to 1
    "default_incoming": 0,        # pending receipts per SKU (units) when not in planning_inputs
    "default_order_multiple": 1,  # round orders up to a multiple of this (e.g. box size)
    "trend_rising": 1.1,          # WH trend = Rising when last window / previous >= this
    "trend_falling": 0.9,         # WH trend = Falling when last window / previous <= this
    "show_export_note": True,     # add 'Latest WH sale ... confirm export' note when data ends before yesterday

    # ---- SKU classes (fast movers get protected, junk does not get ordered) ----
    "abc_cutoffs": (0.70, 0.90),  # A = SKUs making the first 70% of volume, B = up to 90%, C = rest
    "class_reorder_below_days": {"A": 14, "B": 12},   # fast movers are reordered earlier (C uses base value)
    "safety_factor": 1.0,         # base multiplier on demand over (lead + target) days
    "class_safety_factor": {"A": 1.3},                # extra cushion for A items (backtest: ~94% fill rate)
    "min_units_auto": 10,         # Regular SKU needs at least this many units in 90 days to be ordered automatically
    # ---- Open PO (Incoming QTY) ----
    "open_po_qty_factor": 1,      # open PO qty x this = warehouse units (e.g. 12 if the PO sheet counts boxes of 12)
    # ---- New / one-off items: stores decide ----
    "new_item_factor": 0.5,       # expected daily demand of a new item = this x units sold in latest window / window_days
    "new_item_store_periods": 2,  # store sale in any of the latest N store months (incl. MTD) counts as "stores selling"
}
# =========================================================================


def trim(value):
    """Excel TRIM equivalent, also handling tabs/newlines/nonbreaking spaces."""
    return " ".join(str(value).replace("\u00a0", " ").split()) if value is not None else ""


def header_key(value):
    return re.sub(r"[^a-z0-9]", "", trim(value).lower())


# Remove DELETED and common misspellings as separate words, not inside names.
DELETED_WORD = re.compile(r"(?<![A-Za-z])(?:deleted|delted|deletted|deletd|delet|delete|delated|deletedd)(?![A-Za-z])", re.I)


def clean_upc(value):
    if value is None or trim(value) == "":
        return "-"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, float) and not math.isfinite(value):
            return "-"
        value = format(Decimal(str(value)), "f")
    text = trim(DELETED_WORD.sub("", str(value)))
    # A single numeric code gets Excel General -> Text style normalization.
    # Normalize every numeric token in composite cells as well as single codes.
    if re.fullmatch(r"\d+(?:\.0+)?", text):
        text = str(int(text.split(".")[0]))
    elif re.fullmatch(r"\d+(?:\.\d+)?[eE]\+?\d+", text):
        number = Decimal(text)
        if number == number.to_integral_value():
            text = format(number, ".0f")
    text = re.sub(r"(?<![\w.])\d+(?:\.0+)?(?![\w.])",
                  lambda match: str(int(match.group().split('.')[0])), text)
    return text or "-"


# End-only, repeatable suffix stripping. Handles NUPC, N.UPC, N UPC,
# NEW UPC, parenthesized forms, NEW, -NEW, - N and trailing punctuation.
UPC_SUFFIX = r"(?:N\s*[./\-]?\s*UPC|NEW\s*[./\-]?\s*UPC)"
SUFFIX_PATTERNS = [
    re.compile(r"\s*\(\s*(?:" + UPC_SUFFIX + r"|NEW|\.)\s*\)\s*$", re.I),
    re.compile(r"(?:^|[\s\-–—.])(?:" + UPC_SUFFIX + r"|NEW)\s*$", re.I),
    re.compile(r"\s*[-–—]\s*N\s*$", re.I),
    re.compile(r"[\s.\-–—]+$"),
]


def clean_description(value):
    text = trim(value)
    extras = [re.compile(r"(?:^|[\s\-])" + re.escape(s) + r"\s*$", re.I)
              for s in EXTRA_DESCRIPTION_SUFFIXES if s]
    while text:
        previous = text
        for pattern in SUFFIX_PATTERNS + extras:
            text = trim(pattern.sub("", text))
        if text == previous:
            break
    return text


def proper(value):
    """Equivalent to PROPER for the product names in these exports."""
    return trim(value).title()


def description_key(value):
    # Do not strip pack-size/model punctuation or PB product suffixes.
    return trim(value).casefold()


def upc_tokens(value):
    """Split comma/space/newline/semicolon/pipe separated codes for lookup.

    Whole-number UPCs normalize leading zeroes. Text-only identifiers are kept.
    There is no substring or fuzzy UPC match.
    """
    value = clean_upc(value)
    if value == "-":
        return []
    parts = re.split(r"[,;|\s]+", value)
    result = []
    for part in parts:
        code = clean_upc(part.strip('()[]'))
        if code != "-" and code.casefold() not in {x.casefold() for x in result}:
            result.append(code)
    return result


def join_codes(values):
    seen, result = set(), []
    for value in values:
        for code in upc_tokens(value):
            if code.casefold() not in seen:
                result.append(code)
                seen.add(code.casefold())
    return ", ".join(result) or "-"


def number(value, context):
    if value is None or trim(value) == "":
        return None  # blank is not silently replaced with zero
    if isinstance(value, (int, float)):
        return value
    text = trim(value).replace(",", "")
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    try:
        result = Decimal(text)
        if not result.is_finite():
            raise InvalidOperation
        return int(result) if result == result.to_integral_value() else float(result)
    except InvalidOperation:
        raise ValueError(f"Invalid quantity at {context}: {value!r}")


def parse_date(value, context):
    if value is None or trim(value) == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    for pattern in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%m/%d/%Y %H:%M:%S"):
        try:
            return datetime.strptime(trim(value), pattern)
        except ValueError:
            pass
    raise ValueError(f"Unrecognized date at {context}: {value!r}")


def find_source(workbook, kind):
    required = {"inventory": ("productservice", "quantityonhand", "sku"),
                "purchase": ("transactiontype", "qty", "memodescription"),
                "sales": ("transactiontype", "quantity", "salesprice")}[kind]
    name = SHEET_NAMES[kind]
    sheets = [workbook[name]] if name else workbook.worksheets
    for sheet in sheets:
        # Do not rely on max_row, which can be missing in some QB exports.
        for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=20), 1):
            columns = {header_key(c.value): c.column for c in row if c.value is not None}
            if all(key in columns for key in required):
                return sheet, row_number, columns
    raise ValueError(f"Could not find the {kind} headers. Check file/sheet selection.")


def require(columns, *aliases):
    for alias in aliases:
        if header_key(alias) in columns:
            return columns[header_key(alias)]
    raise ValueError(f"Missing required column: {' / '.join(aliases)}")


def copy_raw(source, target):
    """Copy cell content/formulas, styles, merges and relevant sheet layout."""
    style_cache = {}
    for row in source.iter_rows():
        for cell in row:
            if cell.value is None and not cell.has_style:
                continue
            dst = target.cell(cell.row, cell.column, cell.value)
            if cell.has_style:
                # Register each source style once in the destination workbook.
                # Reusing its destination style avoids slow per-cell copying.
                if cell.style_id not in style_cache:
                    dst.font = copy(cell.font)
                    dst.fill = copy(cell.fill)
                    dst.border = copy(cell.border)
                    dst.alignment = copy(cell.alignment)
                    dst.protection = copy(cell.protection)
                    dst.number_format = cell.number_format
                    style_cache[cell.style_id] = copy(dst._style)
                else:
                    dst._style = copy(style_cache[cell.style_id])
            if cell.hyperlink:
                dst._hyperlink = copy(cell.hyperlink)
            if cell.comment:
                dst.comment = copy(cell.comment)
    for key, dim in source.row_dimensions.items():
        target.row_dimensions[key] = copy(dim)
        target.row_dimensions[key].parent = target
    for key, dim in source.column_dimensions.items():
        target.column_dimensions[key] = copy(dim)
        target.column_dimensions[key].parent = target
    for merged in source.merged_cells.ranges:
        target.merge_cells(str(merged))
    target.sheet_format = copy(source.sheet_format)
    target.sheet_properties = copy(source.sheet_properties)
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    target.freeze_panes = source.freeze_panes
    target.auto_filter = copy(source.auto_filter)
    target.page_margins = copy(source.page_margins)
    target.page_setup = copy(source.page_setup)
    target.print_options = copy(source.print_options)
    if source.print_area:
        target.print_area = str(source.print_area).split("!")[-1]
    if source.print_title_rows:
        target.print_title_rows = source.print_title_rows
    if source.print_title_cols:
        target.print_title_cols = source.print_title_cols


def inventory_rows(sheet, header_row, cols, stats):
    # Correct QB mapping: Memo/Description is master UPC; SKU is box UPC.
    box = require(cols, "SKU")
    master = require(cols, "Memo/Description")
    desc = require(cols, "Product/Service")
    qty = require(cols, "Quantity on hand")
    rows = []
    for row_number, cells in enumerate(sheet.iter_rows(min_row=header_row + 1), header_row + 1):
        values = [c.value for c in cells]
        get = lambda n: values[n - 1] if n <= len(values) else None
        original = trim(get(desc))
        if not original and get(qty) is None:
            continue
        if re.search(r"\bGMT[+-]\d", original, re.I) and get(qty) is None:
            stats["footer rows removed"] += 1
            continue
        if original.lower() in ("total", "grand total") and get(box) is None and get(master) is None:
            stats["total rows removed"] += 1
            continue
        quantity = number(get(qty), f"{sheet.title}!row {row_number}")
        cleaned = clean_description(original)
        stats["descriptions changed"] += cleaned != original
        rows.append([clean_upc(get(box)), clean_upc(get(master)), cleaned, quantity])
    if SORT_CLEAN_SHEETS:
        rows.sort(key=lambda r: (r[2].casefold(), r[0], r[1]))
    return ["box upc", "master upc", "description", "QTY"], rows


def transaction_rows(sheet, header_row, cols, kind, stats):
    dtype = require(cols, "Transaction Type")
    ddate = require(cols, "Date", "Transaction date")
    invoice = require(cols, "Num")
    desc = require(cols, "Product/Service", "Product/Service full name")
    upc = require(cols, "Memo/Description", "Description")
    qty = require(cols, "Qty", "Quantity")
    wanted = "bill" if kind == "purchase" else "invoice"
    rate_col = next((cols[k] for k in ("rate", "unitprice", "cost", "price") if k in cols), None) if kind == "purchase" else None
    if kind == "purchase" and rate_col is None:
        print("  WARNING: no Rate column in the purchase export; buying prices will stay blank.", flush=True)
    rows, store = [], ""
    for row_number, cells in enumerate(sheet.iter_rows(min_row=header_row + 1), header_row + 1):
        values = [c.value for c in cells]
        get = lambda n: values[n - 1] if n <= len(values) else None
        group = trim(get(1))
        if group and not get(ddate) and not get(dtype):
            # Store headings are kept only in sales; totals aren't store names.
            if not re.match(r"^(?:total\b|grand total\b)", group, re.I) and not re.search(r"GMT[+-]\d", group):
                store = group
        if trim(get(dtype)).casefold() != wanted:
            if get(dtype):
                stats["nonmatching transaction rows removed"] += 1
            continue
        original = trim(get(desc))
        if kind in ("purchase", "sales"):
            # QB full-name paths may contain several category levels.
            # Keep only the product name after the final colon, then apply
            # the existing suffix cleanup. Original paths stay in purchase raw.
            original = trim(original.rsplit(":", 1)[-1])
        cleaned = clean_description(original)
        stats["descriptions changed"] += cleaned != original
        dt = parse_date(get(ddate), f"{sheet.title}!row {row_number}")
        num = get(invoice)
        if isinstance(num, (int, float)) and float(num).is_integer():
            num = str(int(num))
        else:
            num = trim(num)
        row = [dt, trim(get(dtype)), num, cleaned, clean_upc(get(upc)),
               number(get(qty), f"{sheet.title}!row {row_number}")]
        if kind == "sales":
            row.insert(0, store)
            if not store:
                stats["invoice rows missing store name"] += 1
        if kind == "purchase":
            # The group heading above the bill lines is the vendor. Rate = unit buying price.
            row += [store, number(get(rate_col), f"{sheet.title}!row {row_number}") if rate_col else None]
        rows.append(row)
    if SORT_CLEAN_SHEETS:
        if kind == "sales":
            rows.sort(key=lambda r: (r[0].casefold(), r[1] or datetime.min, r[3], r[4].casefold()))
        else:
            rows.sort(key=lambda r: (r[0] or datetime.min, r[2], r[3].casefold()))
    headers = ["Date", "Transaction Type", "Num", "description", "master upc", "QTY"]
    if kind == "sales":
        headers[0] = "Transaction date"
        headers.insert(0, "Store name")
    if kind == "purchase":
        headers += ["Vendor", "Rate"]
    return headers, rows


def read_pricebook(path):
    """Import only the authoritative BTCM table at E:I in PB.xlsx.

    A:C and other sheets are excluded from both output and lookup indexes.
    Source sheet/row/block remain visible in output.
    """
    records = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in (workbook[PRICEBOOK_SHEET],):
            header = None
            for cells in sheet.iter_rows(min_row=1, max_row=20, values_only=True):
                keys = [header_key(v) for v in cells]
                if any(k in ("description", "desc", "discription") for k in keys) and any(k in ("upc", "masterupc", "boxupc") for k in keys):
                    header = (cells, keys)
                    break
            if header is None:
                continue
            # These QB/MPB exports put headers on row 1; find the detected row.
            for row_number, cells in enumerate(sheet.iter_rows(min_row=1, max_row=20, values_only=True), 1):
                if cells == header[0]:
                    header_row = row_number
                    break
            keys = header[1]
            if len(keys) < 9 or keys[4] != "boxupc" or keys[5] not in ("upc", "masterupc") or keys[6] not in ("desc", "description", "discription"):
                raise ValueError("PB.xlsx requires BOX UPC in E, master UPC in F, description in G.")
            desc_cols = [i for i, k in enumerate(keys) if k in ("description", "desc", "discription")]
            blocks = []
            for desc in desc_cols:
                # Import only column G (the authoritative E:I table).
                if desc != 6:
                    continue
                left = max((d + 1 for d in desc_cols if d < desc), default=0)
                right = min((d for d in desc_cols if d > desc), default=len(keys))
                master_candidates = [i for i in range(left, desc) if keys[i] in ("masterupc", "upc", "productupc")]
                if not master_candidates:
                    continue
                master = master_candidates[-1]
                boxes = [i for i in range(left, desc) if keys[i] == "boxupc"]
                box = boxes[-1] if boxes else None
                # Following metadata belongs to this block until the next UPC.
                end = min((i for i in range(desc + 1, right) if keys[i] in ("boxupc", "upc", "masterupc")), default=right)
                attr = lambda aliases: next((i for i in range(master + 1, end) if keys[i] in aliases), None)
                blocks.append((box, master, desc, attr(("brand",)), attr(("department", "depart", "departfinal")), attr(("packof", "pack")), 0 if box is not None or keys[master] == "masterupc" else 1))
            for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
                for block_number, (box, master, desc, brand, dept, pack, priority) in enumerate(blocks, 1):
                    get = lambda i: row[i] if i is not None and i < len(row) else None
                    name = trim(get(desc))
                    if not name or name.startswith("#"):
                        continue
                    records.append({"box": clean_upc(get(box)), "master": clean_upc(get(master)),
                                    "description": name, "brand": trim(get(brand)),
                                    "department": trim(get(dept)), "pack": trim(get(pack)),
                                    "sheet": sheet.title, "row": row_number,
                                    "block": 2, "priority": priority})
    finally:
        workbook.close()
    if not records:
        raise ValueError("No Price Book UPC/description blocks found. Check headers.")
    return records


def pricebook_indexes(records):
    indexes = {"master": defaultdict(list), "box": defaultdict(list), "description": defaultdict(list)}
    for record in records:
        for field in ("master", "box"):
            for code in upc_tokens(record[field]):
                indexes[field][code.casefold()].append(record)
        indexes["description"][description_key(record["description"])].append(record)
    return indexes


def resolve_candidates(candidates, qb_description):
    if not candidates:
        return None
    best_priority = min(r["priority"] for r in candidates)
    candidates = [r for r in candidates if r["priority"] == best_priority]
    distinct = {}
    for record in candidates:
        distinct.setdefault(description_key(record["description"]), record)
    def attrs(records):
        out = []
        for field in ("department", "brand", "pack"):
            seen = []
            for r in records:
                v = trim(r.get(field))
                if v.startswith("#"):  # '#N/A' etc. from the price book: leave blank
                    continue
                if v and v.casefold() not in {x.casefold() for x in seen}:
                    seen.append(v)
            out.append(" | ".join(seen))
        return tuple(out)
    if len(distinct) == 1:
        record = next(iter(distinct.values()))
        return "Matched", record["description"], attrs([record])
    exact = distinct.get(description_key(qb_description))
    if exact:
        return "Matched", exact["description"], attrs([exact])
    return "Ambiguous - review", " | ".join(r["description"] for r in distinct.values()), attrs(list(distinct.values()))


def lookup_inventory(rows, indexes):
    output = []
    counts = Counter()
    for box, master, qb_description, qty in rows:
        result, method, matched_code = None, "Not found", "-"
        # Stages 1-3 use only single-code cells. Multi-code cells are deferred
        # until after the description lookup, exactly as requested.
        rules = [(master, "master", "Master to master"),
                 (master, "box", "Master to box"),
                 (box, "master", "Box to master")]
        for value, index_name, label in rules:
            tokens = upc_tokens(value)
            if len(tokens) != 1:
                continue
            for code in tokens:
                result = resolve_candidates(indexes[index_name].get(code.casefold(), []), qb_description)
                if result:
                    method, matched_code = label, code
                    break
            if result:
                break
        if result is None and qb_description:
            result = resolve_candidates(indexes["description"].get(description_key(qb_description), []), qb_description)
            if result:
                method = "Description to description"
        if result is None:
            # Stage 5: all multi-code cells against PB master UPC first,
            # then PB box UPC. Inventory master cell precedes box cell.
            for index_name in ("master", "box"):
                for value, field_name in ((master, "master"), (box, "box")):
                    tokens = upc_tokens(value)
                    if len(tokens) < 2:
                        continue
                    for code in tokens:
                        result = resolve_candidates(indexes[index_name].get(code.casefold(), []), qb_description)
                        if result:
                            method = f"Multiple {field_name} UPCs to PB {index_name}"
                            matched_code = code
                            break
                    if result:
                        break
                if result:
                    break
        status, name, pb_attrs = result if result else ("Not found", "-", ("", "", ""))
        counts[status] += 1
        if method != status:
            counts[method] += 1
        output.append([box, master, qb_description, qty, name, method, matched_code, status, *pb_attrs])
    return output, counts


def summarize_inventory(rows):
    groups = {}
    for row_number, row in enumerate(rows, 2):
        box, master, qb_description, qty, pb_description, method, matched_code, status = row[:8]
        pb_attrs = row[8:11] if len(row) >= 11 else ("", "", "")
        # User's ordering identity is the cleaned QB description, not PB name.
        key = ("qb", description_key(qb_description)) if trim(qb_description) else ("row", row_number)
        if key not in groups:
            groups[key] = {"box": [], "master": [], "qb": [], "pb": [],
                           "qty": Decimal(0), "has_qty": False, "methods": [], "status": status, "count": 0,
                           "attrs": [[], [], []]}
        group = groups[key]
        group["box"].append(box)
        group["master"].append(master)
        for bucket, value in zip(group["attrs"], pb_attrs):
            if value and value.casefold() not in {x.casefold() for x in bucket}:
                bucket.append(value)
        for field, value in (("qb", qb_description), ("methods", method)):
            if value and value.casefold() not in {x.casefold() for x in group[field]}:
                group[field].append(value)
        for name in pb_description.split(" | "):
            if name and name != "-" and name.casefold() not in {x.casefold() for x in group["pb"]}:
                group["pb"].append(name)
        if qty is not None:
            group["qty"] += Decimal(str(qty))
            group["has_qty"] = True
        group["count"] += 1
    output = []
    for group in groups.values():
        q = group["qty"]
        qty = (int(q) if q == q.to_integral_value() else float(q)) if group["has_qty"] else None
        output.append([join_codes(group["box"]), join_codes(group["master"]),
                       " | ".join(group["qb"]), qty, " | ".join(group["pb"]) or "-", group["status"],
                       group["count"], " | ".join(group["methods"]),
                       *[" | ".join(b) for b in group["attrs"]]])
    if SORT_CLEAN_SHEETS:
        output.sort(key=lambda r: (r[2].casefold(), r[0]))
    before = sum((Decimal(str(r[3])) for r in rows if r[3] is not None), Decimal(0))
    after = sum((Decimal(str(r[3])) for r in output if r[3] is not None), Decimal(0))
    if abs(before - after) > Decimal("0.000000001"):
        raise ValueError(f"Inventory quantity reconciliation failed: {before} vs {after}")
    return output


def reverse_purchase_check(summary_rows, purchase_rows, pb_indexes):
    """Match purchases back to summarized inventory with exact indexed keys.

    Returns descriptions/0 for purchase clean and matched inventory row indexes.
    Description first, then all master/box UPC tokens, PB names and PB mapping.
    No substring/fuzzy matching. Multiple exact candidates remain visible.
    """
    by_name, by_master, by_box, by_pb = (defaultdict(set) for _ in range(4))
    def name_keys(value):
        # QB purchase paths may include a category, e.g. Display:Product Name.
        # Match the exact leaf name too; never use fuzzy/partial word matches.
        full = description_key(value)
        leaf = description_key(trim(value).rsplit(":", 1)[-1])
        return list(dict.fromkeys(key for key in (full, leaf) if key))
    for index, row in enumerate(summary_rows):
        for key in name_keys(row[2]):
            by_name[key].add(index)
        for value, lookup in ((row[1], by_master), (row[0], by_box)):
            for token in upc_tokens(value):
                lookup[token.casefold()].add(index)
        for name in row[4].split(" | "):
            if name != "-" and trim(name):
                for key in name_keys(name):
                    by_pb[key].add(index)
    checks, matched_indexes = [], set()
    for purchase in purchase_rows:
        name, upc = purchase[3], purchase[4]
        candidates = set()
        for key in name_keys(name):
            candidates.update(by_name.get(key, set()))
            if candidates:
                break
        if not candidates:
            for lookup in (by_master, by_box):
                for token in upc_tokens(upc):
                    candidates.update(lookup.get(token.casefold(), set()))
                if candidates:
                    break
        if not candidates and trim(name):
            for key in name_keys(name):
                candidates.update(by_pb.get(key, set()))
                if candidates:
                    break
        if not candidates:
            # Changed purchase UPC may still map through the authoritative PB.
            for field in ("master", "box"):
                for token in upc_tokens(upc):
                    records = pb_indexes[field].get(token.casefold(), [])
                    for record in records:
                        pb_name = description_key(record["description"])
                        candidates.update(by_pb.get(pb_name, set()))
                        candidates.update(by_name.get(pb_name, set()))
                if candidates:
                    break
        if not candidates and trim(name):
            for record in pb_indexes["description"].get(description_key(name), []):
                for value in (record["master"], record["box"]):
                    for token in upc_tokens(value):
                        candidates.update(by_master.get(token.casefold(), set()))
                        candidates.update(by_box.get(token.casefold(), set()))
        descriptions = []
        for index in sorted(candidates):
            description = summary_rows[index][2]
            if description and description not in descriptions:
                descriptions.append(description)
        checks.append(" | ".join(descriptions) if descriptions else 0)
        matched_indexes.update(candidates)
    return checks, matched_indexes


def purchase_activity(summary_rows, purchase_rows, reverse_matches=None, force_active=None):
    """Any cleaned Bill item match is Active, irrespective of quantity.

    Purchase export has master UPC only; never invent a purchase box field.
    Single UPC matching precedes description matching; composite cells follow.
    PB descriptions are used only for exact description fallback, not quantities.
    """
    codes = set()
    names = set()
    for row in purchase_rows:
        # Date, Transaction Type, Num, description, master upc, QTY
        if trim(row[1]).casefold() != "bill":
            continue
        codes.update(code.casefold() for code in upc_tokens(row[4]))
        if trim(row[3]):
            names.add(description_key(row[3]))
    audit, active = [], []
    for summary_index, row in enumerate(summary_rows):
        base = list(row[:5])
        extra = list(row[8:11]) if len(row) >= 11 else ["", "", ""]
        box, master, qb_description, qty, pb_description = base
        method, matched_value = "Not found", "-"
        for value, label in ((master, "Master to purchase master"), (box, "Box to purchase master")):
            tokens = upc_tokens(value)
            if len(tokens) == 1 and tokens[0].casefold() in codes:
                method, matched_value = label, tokens[0]
                break
        if method == "Not found" and trim(qb_description) and description_key(qb_description) in names:
            method, matched_value = "Inventory description to purchase description", qb_description
        if method == "Not found":
            for name in pb_description.split(" | "):
                if name != "-" and trim(name) and description_key(name) in names:
                    method, matched_value = "PB description to purchase description", name
                    break
        if method == "Not found":
            for value, label in ((master, "Multiple master UPCs to purchase master"), (box, "Multiple box UPCs to purchase master")):
                tokens = upc_tokens(value)
                if len(tokens) < 2:
                    continue
                for code in tokens:
                    if code.casefold() in codes:
                        method, matched_value = label, code
                        break
                if method != "Not found":
                    break
        if method == "Not found" and reverse_matches and summary_index in reverse_matches:
            method, matched_value = "Purchase reverse lookup", qb_description
        status = "Inactive" if method == "Not found" else "Active"
        if status == "Inactive" and force_active and summary_index in force_active:
            # Re-ordered item: it has an open PO now, so it is planned like any other active SKU.
            status, method, matched_value = OPEN_PO_STATUS, "Open PO", force_active[summary_index]
        audit.append(base + [status, method, matched_value])
        if status in ("Active", OPEN_PO_STATUS):
            active.append(base + [status] + extra)
    return audit, active


def write_clean(sheet, headers, rows):
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for c in sheet[1]:
        c.fill = PatternFill("solid", fgColor="17365D")
        c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 25
    for i, name in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(i)].width = (
            58 if "description" in name.lower() else 42 if name == "Store name"
            else 32 if "upc" in name else 22 if "date" in name.lower()
            else 20 if name == "Transaction Type" else 16)
        fmt = "@" if "upc" in name.lower() or name == "Num" else "mm/dd/yyyy" if "date" in name.lower() else "General"
        if fmt != "General":
            for cells in sheet.iter_rows(min_row=2, min_col=i, max_col=i):
                cells[0].number_format = fmt


def append_date_range(sheet, kind, headers, rows):
    start, end = APPEND_COVERAGE[kind]
    if start or end:
        if not start or not end:
            raise ValueError("Both append coverage dates are required.")
        return datetime.strptime(start, "%Y-%m-%d").date(), datetime.strptime(end, "%Y-%m-%d").date()
    label = trim(sheet.cell(3, 1).value)
    for pattern in ("%B %d, %Y", "%B %d %Y", "%m/%d/%Y", "%Y-%m-%d"):
        try:
            day = datetime.strptime(label, pattern).date()
            return day, day
        except ValueError:
            pass
    match = re.fullmatch(r"([A-Za-z]+)\s+(\d{1,2})\s*-\s*(\d{1,2}),?\s+(\d{4})", label)
    if match:
        month = datetime.strptime(match[1], "%B").month
        return date(int(match[4]), month, int(match[2])), date(int(match[4]), month, int(match[3]))
    for pattern in ("%B %Y", "%B, %Y"):
        try:
            first = datetime.strptime(label, pattern).date()
            return first, date(first.year, first.month, calendar.monthrange(first.year, first.month)[1])
        except ValueError:
            pass
    # Read dates on all raw transactions, including non-Bills/non-Invoices.
    # This supports daily reports with zero qualifying items.
    date_col = require(headers, "Date", "Transaction date")
    dates = []
    for values in sheet.iter_rows(min_row=6, values_only=True):
        value = values[date_col - 1] if len(values) >= date_col else None
        if value:
            try:
                parsed = parse_date(value, "append report")
                if parsed:
                    dates.append(parsed.date())
            except ValueError:
                pass
    if not dates:
        raise ValueError(f"Cannot determine {kind} append dates. Set APPEND_COVERAGE explicitly.")
    print("  Report period not recognized; using complete-export transaction date range.", flush=True)
    return min(dates), max(dates)


def merge_daily_history(base_rows, append_rows, kind, history, coverage=None):
    """Replace covered dates, not identical lines: repeated invoice items survive.

    Stored daily overrides persist when tomorrow's file replaces today's file.
    Reimporting the same export does not duplicate any quantities.
    """
    overrides = history.setdefault(kind, {})
    date_index = 0 if kind == "purchase" else 1
    if coverage:
        start, end = coverage
        if end < start:
            raise ValueError("Append end date cannot precede start date.")
        if (end - start).days > 370:
            raise ValueError("Append coverage exceeds 370 days; check dates.")
        grouped = defaultdict(list)
        for row in append_rows:
            if row[date_index] is None:
                raise ValueError("Append transaction has a blank date.")
            day = row[date_index].date()
            if not start <= day <= end:
                raise ValueError("Append transaction is outside declared coverage.")
            serialized = list(row)
            serialized[date_index] = row[date_index].isoformat()
            grouped[day.isoformat()].append(serialized)
        day = start
        while day <= end:
            overrides[day.isoformat()] = grouped.get(day.isoformat(), [])
            day += timedelta(days=1)
    combined = [r for r in base_rows if r[date_index] is None or r[date_index].date().isoformat() not in overrides]
    for day, stored_rows in sorted(overrides.items()):
        for row in stored_rows:
            restored = list(row)
            restored[date_index] = datetime.fromisoformat(restored[date_index])
            # Apply the latest normalization rules to persisted append history.
            desc_index, upc_index = (3, 4) if kind == "purchase" else (4, 5)
            restored[desc_index] = clean_description(trim(restored[desc_index]).rsplit(":", 1)[-1])
            restored[upc_index] = clean_upc(restored[upc_index])
            if kind == "purchase" and len(restored) < 8:
                restored += [""] * (7 - len(restored)) + [None]  # history saved before Vendor/Rate existed
            combined.append(restored)
    if SORT_CLEAN_SHEETS:
        if kind == "purchase":
            combined.sort(key=lambda r: (r[0] or datetime.min, r[2], r[3].casefold()))
        else:
            combined.sort(key=lambda r: (r[0].casefold(), r[1] or datetime.min, r[3], r[4].casefold()))
    return combined


def add_lookup_columns(sheet, summarized_checks, active_checks):
    for values, header in ((summarized_checks, "lookup inv summarized"), (active_checks, "lookup active inventory")):
        column = sheet.max_column + 1
        sheet.cell(1, column, header)
        sheet.cell(1, column)._style = copy(sheet.cell(1, 4)._style)
        sheet.column_dimensions[get_column_letter(column)].width = 58
        for number, value in enumerate(values, 2):
            sheet.cell(number, column, value)
    sheet.auto_filter.ref = sheet.dimensions


def month_sequence(start, end):
    result = []
    total = start.year * 12 + start.month - 1
    final = end.year * 12 + end.month - 1
    while total <= final:
        year, month0 = divmod(total, 12)
        result.append((year, month0 + 1))
        total += 1
    return result


def sales_coverage_months(sheet, sales_rows, history):
    covered = set()
    start, end = SALES_BASE_COVERAGE
    if start and end:
        covered.update(month_sequence(date.fromisoformat(start), date.fromisoformat(end)))
    else:
        label = trim(sheet.cell(3, 1).value)
        match = re.fullmatch(r"([A-Za-z]+)\s*-\s*([A-Za-z]+),?\s+(\d{4})", label)
        if match:
            first = datetime.strptime(match[1], "%B").month
            last = datetime.strptime(match[2], "%B").month
            year = int(match[3])
            covered.update(month_sequence(date(year if first <= last else year-1, first, 1), date(year,last,1)))
        else:
            for pattern in ("%B %Y", "%B, %Y"):
                try:
                    parsed = datetime.strptime(label, pattern)
                    covered.add((parsed.year, parsed.month))
                    break
                except ValueError:
                    pass
    # Actual transactions/complete daily imports confirm at least some coverage.
    covered.update((r[1].year, r[1].month) for r in sales_rows if r[1] is not None)
    covered.update((date.fromisoformat(day).year, date.fromisoformat(day).month) for day in history.get("sales", {}))
    return covered


def monthly_sales_rows(active_rows, sales_rows, checks, covered, as_of):
    current = (as_of.year, as_of.month)
    total = as_of.year * 12 + as_of.month - 1
    year, month0 = divmod(total - 6, 12)
    months = month_sequence(date(year, month0 + 1, 1), as_of)
    by_name = {description_key(row[2]): i for i, row in enumerate(active_rows)}
    quantities = defaultdict(Decimal)
    audit = Counter()
    for row, check in zip(sales_rows, checks):
        transaction_date, qty = row[1], row[6]
        if transaction_date is None or transaction_date.date() > as_of:
            audit["Missing/future date rows"] += 1
            continue
        month = (transaction_date.year, transaction_date.month)
        if month not in months:
            audit["Rows outside sales window"] += 1
            continue
        if qty is None:
            audit["Rows with missing quantity"] += 1
            continue
        audit["Window sales quantity"] += Decimal(str(qty))
        if check == 0:
            audit["Unmatched/inactive sales quantity"] += Decimal(str(qty))
            continue
        candidates = {by_name[key] for key in (description_key(name) for name in str(check).split(" | ")) if key in by_name}
        if len(candidates) != 1:
            audit["Ambiguous sales quantity"] += Decimal(str(qty))
            audit["Ambiguous sales rows"] += 1
            continue
        index = next(iter(candidates))
        quantities[(index, month)] += Decimal(str(qty))
        audit["Assigned active sales quantity"] += Decimal(str(qty))
    if audit["Window sales quantity"] != audit["Assigned active sales quantity"] + audit["Unmatched/inactive sales quantity"] + audit["Ambiguous sales quantity"]:
        raise ValueError("Monthly sales quantity reconciliation failed.")
    headers = list(FINAL_BASE_HEADERS)
    headers.extend("Warehouse " + date(y,m,1).strftime("%b %Y") + (" MTD" if (y,m)==current else "") for y,m in months)
    output = []
    for index, active in enumerate(active_rows):
        values = []
        for month in months:
            value = quantities[(index,month)]
            values.append((int(value) if value == value.to_integral_value() else float(value)) if month in covered else "Missing")
        box, master, desc, qty_value, pb_desc, status = active[:6]
        dept, brand, pack = (list(active[6:9]) + ["", "", ""])[:3]
        flag = ""
        if isinstance(qty_value, (int, float)) and math.isfinite(float(qty_value)) and qty_value < 0:
            # Negative stock is treated as 0 in the final sheet (and in the planning); the flag keeps the original.
            flag = f"Was negative ({qty_value:g})"
            qty_value = 0
        output.append([status, dept, brand, box, master, desc, pack, qty_value, flag, pb_desc] + values)
    return headers, output, audit


def store_month(value, default_year, context):
    if isinstance(value, (date, datetime)):
        return date(value.year, value.month, 1)
    label = trim(value)
    for pattern in ("%B %Y", "%b %Y", "%Y-%m", "%B-%Y", "%b-%Y", "%m/%Y"):
        try:
            d = datetime.strptime(label, pattern)
            return date(d.year, d.month, 1)
        except ValueError:
            pass
    for pattern in ("%B", "%b"):
        try:
            return date(default_year, datetime.strptime(label, pattern).month, 1)
        except ValueError:
            pass
    raise ValueError(f"Invalid store Month at {context}: {value!r}")


def read_store_sales(path, target, raw_prefix, default_year):
    """Rows: month ISO, optional day ISO, description, scan UPC, units, amount.

    Retail is ignored entirely. Every tab with the expected headers is included.
    Month-only data is a monthly snapshot; dated data supports daily replacement.
    """
    source = load_workbook(path, read_only=True, data_only=True)
    result, covered, included = [], set(), []
    modes = set()
    try:
        for sheet in source:
            if re.search(r"pivot|povit|\bpvt\b", sheet.title, re.I):
                continue
            found = None
            for rn, row in enumerate(sheet.iter_rows(max_row=20, values_only=True), 1):
                cols = {header_key(v): i for i, v in enumerate(row) if v is not None}
                if all(k in cols for k in ("scancode", "description", "totalretail")):
                    found = rn, cols
                    break
            if found is None:
                continue
            rn, cols = found
            date_col = next((cols[k] for k in ("date", "transactiondate", "salesdate") if k in cols), None)
            if "month" not in cols and date_col is None:
                # A named monthly tab can supply Month if its source header omits it.
                try:
                    named_month = store_month(sheet.title, default_year, str(path))
                except ValueError:
                    raise ValueError(f"Store tab {sheet.title!r} needs a Month or Date column.")
            else:
                named_month = None
            included.append(sheet.title)
            name = raw_prefix if len(included) == 1 else raw_prefix[:27] + f" {len(included)}"
            raw = target.create_sheet(name[:31])
            # Preserve source values, including optional Retail, without importing formatting.
            for original in sheet.iter_rows(values_only=True):
                raw.append(original)
            for rowno, row in enumerate(sheet.iter_rows(min_row=rn+1, values_only=True), rn+1):
                def cell(index):
                    return row[index] if index is not None and index < len(row) else None
                code, desc, amount = cell(cols["scancode"]), cell(cols["description"]), cell(cols["totalretail"])
                if all(v is None or trim(v) == "" for v in (code, desc, amount)):
                    continue
                if header_key(code) == "scancode" and header_key(desc) == "description":
                    continue
                context = f"{path.name}/{sheet.title}/row {rowno}"
                if trim(desc).casefold() in ("total", "grand total") and not trim(code):
                    continue
                day = parse_date(cell(date_col), context) if date_col is not None else None
                if date_col is not None and day is None:
                    raise ValueError(f"Blank Date in dated store export at {context}")
                month_value = cell(cols.get("month"))
                if month_value is not None and trim(month_value):
                    month = store_month(month_value, default_year, context)
                    if day and (day.year, day.month) != (month.year, month.month):
                        raise ValueError(f"Date/Month disagreement at {context}")
                else:
                    month = date(day.year, day.month, 1) if day else named_month
                if month is None:
                    raise ValueError(f"Blank Month at {context}")
                amount = number(trim(amount).replace("$", ""), context + " Total Retail")
                if amount is None or not math.isfinite(float(amount)):
                    raise ValueError(f"Missing/invalid store amount at {context}; cannot silently use zero")
                qty = number(cell(cols.get("qtysold")), context + " Qty Sold")
                result.append([month.isoformat(), day.date().isoformat() if day else None,
                               clean_description(trim(desc).rsplit(":", 1)[-1]), clean_upc(code), qty, amount])
                covered.add((month.year, month.month))
                modes.add("daily" if day else "monthly")
        if not included:
            raise ValueError(f"No store sales table found in {path}")
        if len(modes) > 1:
            raise ValueError("Store file mixes dated rows and month-only rows; use one format per file.")
        return result, covered, next(iter(modes), "empty")
    finally:
        source.close()


def merge_store_history(base, appended, append_mode, history):
    """Month snapshots replace months; dated daily exports replace covered days."""
    saved = history.setdefault("stores_sales", {"months": {}, "days": {}})
    monthly, daily = saved["months"], saved["days"]
    if append_mode == "monthly":
        grouped = defaultdict(list)
        for row in appended:
            grouped[row[0]].append(row)
        if not grouped:
            raise ValueError("Empty month-only store append: cannot infer which month to replace.")
        for month, rows in grouped.items():
            monthly[month] = rows
            for day in list(daily):
                if day[:7] == month[:7]:
                    del daily[day]
        print("  Store append: replacing complete MTD/month snapshot(s): " + ", ".join(grouped), flush=True)
    elif append_mode == "daily":
        days = [date.fromisoformat(r[1]) for r in appended]
        start_text, end_text = STORES_APPEND_COVERAGE
        if bool(start_text) != bool(end_text):
            raise ValueError("Set both store append coverage dates.")
        if start_text:
            start, end = date.fromisoformat(start_text), date.fromisoformat(end_text)
        elif days:
            start, end = min(days), max(days)
            print("  Store dated append assumes a complete all-store report for its date range.", flush=True)
        else:
            raise ValueError("Empty dated store append requires STORES_APPEND_COVERAGE.")
        if end < start or (end-start).days > 370:
            raise ValueError("Invalid store append date range.")
        grouped = defaultdict(list)
        base_months = {r[0] for r in base if r[1] is None}
        day = start
        while day <= end:
            month = day.replace(day=1).isoformat()
            if month in monthly or month in base_months:
                raise ValueError("Daily store append overlaps a month-only snapshot. Use a complete MTD export for that month instead.")
            day += timedelta(days=1)
        for row in appended:
            if not start <= date.fromisoformat(row[1]) <= end:
                raise ValueError("Store append row outside declared coverage.")
            grouped[row[1]].append(row)
        day = start
        while day <= end:
            daily[day.isoformat()] = grouped.get(day.isoformat(), [])
            day += timedelta(days=1)
    combined = [r for r in base if r[0] not in monthly and r[1] not in daily]
    for rows in monthly.values():
        combined.extend(rows)
    for rows in daily.values():
        combined.extend(rows)
    combined = [list(r) for r in combined]
    for row in combined:
        row[2], row[3] = clean_description(row[2]), clean_upc(row[3])
    combined.sort(key=lambda r: (r[0], description_key(r[2]), r[3], r[1] or ""))
    return combined


def store_monthly_values(active_rows, rows, checks, covered, as_of):
    total = as_of.year * 12 + as_of.month - 1
    year, month0 = divmod(total-6, 12)
    months = month_sequence(date(year, month0+1, 1), as_of)
    names = {description_key(r[2]): i for i, r in enumerate(active_rows)}
    amounts, audit = defaultdict(Decimal), Counter()
    for row, check in zip(rows, checks):
        month_date = date.fromisoformat(row[0])
        month = month_date.year, month_date.month
        if month not in months or (row[1] and date.fromisoformat(row[1]) > as_of):
            audit["Store rows outside window/future"] += 1
            continue
        amount = Decimal(str(row[5]))
        audit["Store window amount"] += amount
        candidates = {names[k] for k in (description_key(n) for n in str(check).split(" | ")) if k in names} if check != 0 else set()
        if not candidates:
            audit["Store unmatched/inactive amount"] += amount
        elif len(candidates) != 1:
            audit["Store ambiguous amount"] += amount
            audit["Store ambiguous rows"] += 1
        else:
            amounts[next(iter(candidates)), month] += amount
            audit["Store assigned active amount"] += amount
    assert audit["Store window amount"] == audit["Store assigned active amount"] + audit["Store unmatched/inactive amount"] + audit["Store ambiguous amount"]
    headers = ["Stores Amount " + date(y,m,1).strftime("%b %Y") + (" MTD" if (y,m)==(as_of.year,as_of.month) else "") for y,m in months]
    values = [[float(amounts[i, month]) if month in covered else "Missing" for month in months] for i in range(len(active_rows))]
    return headers, values, audit


def store_inventory_presence(active_rows, store_rows, summary_checks, as_of, active_checks=None):
    """All exact presence evidence, independent of unique amount allocation.

    Yellow means no matching row in the report window. It does NOT mean zero
    revenue, or an unresolved split between multiple inventory descriptions.
    Check every UPC token against BOTH inventory code columns, then exact QB/PB
    names and all names returned by the existing summarized-inventory lookup.
    """
    by_code, by_name, by_inventory_name = (defaultdict(set) for _ in range(3))
    for index, row in enumerate(active_rows):
        for field in row[:2]:
            for code in upc_tokens(field):
                by_code[code.casefold()].add(index)
        if trim(row[2]):
            by_inventory_name[description_key(row[2])].add(index)
        for name in [row[2]] + str(row[4] or "").split(" | "):
            if trim(name) and trim(name) != "-":
                by_name[description_key(name)].add(index)
    total = as_of.year * 12 + as_of.month - 1
    year, month0 = divmod(total - 6, 12)
    report_months = set(month_sequence(date(year, month0 + 1, 1), as_of))
    found = set()
    for row_number, (row, summary_check) in enumerate(zip(store_rows, summary_checks)):
        month = date.fromisoformat(row[0])
        if (month.year, month.month) not in report_months:
            continue
        if row[1] and date.fromisoformat(row[1]) > as_of:
            continue
        for code in upc_tokens(row[3]):
            found.update(by_code.get(code.casefold(), ()))
        if trim(row[2]):
            found.update(by_name.get(description_key(row[2]), ()))
        evidence = [summary_check]
        if active_checks is not None:
            evidence.append(active_checks[row_number])
        for check in evidence:
            if check != 0:
                for name in str(check).split(" | "):
                    found.update(by_inventory_name.get(description_key(name), ()))
    return found


def add_stores_reports(paths, target, summarized, active_rows, indexes, history, as_of, warehouse_headers, warehouse_rows):
    output = Path(paths["output"]).expanduser().resolve()
    base_path = Path(paths["stores_sales"]).expanduser().resolve()
    history_path = Path(paths.get("history") or str(output) + ".history.json").expanduser().resolve()
    if base_path in (output, history_path):
        raise ValueError("Store input must differ from output/history.")
    base, covered, mode = read_store_sales(base_path, target, "stores sales raw", STORES_BASE_YEAR)
    appended, append_mode = [], None
    append_name = paths.get("stores_sales_append", "")
    if append_name:
        append_path = Path(append_name).expanduser().resolve()
        if append_path in (base_path, output, history_path):
            raise ValueError("Store append must differ from store base/output.")
        appended, extra, append_mode = read_store_sales(append_path, target, "stores sales append raw", STORES_APPEND_YEAR)
        if append_mode == "empty":
            if all(STORES_APPEND_COVERAGE):
                append_mode = "daily"
            else:
                raise ValueError("Store append is empty; declare dated coverage for a zero-sales day.")
        covered.update(extra)
    rows = merge_store_history(base, appended, append_mode, history)
    covered.update((date.fromisoformat(r[0]).year, date.fromisoformat(r[0]).month) for r in rows)
    saved = history.get("stores_sales", {})
    covered.update((date.fromisoformat(d).year, date.fromisoformat(d).month) for d in saved.get("days", {}))
    # Reuse the EXACT purchase matching engine, converting to its item field layout.
    items = [[None, "Store sale", "", r[2], r[3], r[4]] for r in rows]
    summary_checks, _ = reverse_purchase_check(summarized, items, indexes)
    active_checks, _ = reverse_purchase_check(active_rows, items, indexes)
    clean = target.create_sheet("stores sales clean")
    display = [[datetime.fromisoformat(r[1]) if r[1] else None,
                date.fromisoformat(r[0]).strftime("%B %Y"), r[3], r[2], r[4], r[5]] for r in rows]
    write_clean(clean, ["Date", "Month", "master upc", "description", "Qty Sold (single units)", "Stores Sales Amount"], display)
    add_lookup_columns(clean, summary_checks, active_checks)
    for cells in clean.iter_rows(min_row=2, min_col=6, max_col=6):
        cells[0].number_format = '#,##0.00'
    headers, values, audit = store_monthly_values(active_rows, rows, active_checks, covered, as_of)
    final = target.create_sheet("final sheet")
    write_clean(final, warehouse_headers + headers,
                [list(w) + v for w,v in zip(warehouse_rows, values)])
    for column in range(BASE_COLS+1, final.max_column+1):
        final.column_dimensions[get_column_letter(column)].width = 29
    # Presence and amount allocation are different checks. A multi-UPC sales
    # record proves both inventory codes occur in stores, even when its combined
    # amount cannot be assigned to one inventory row without double counting.
    matched_active = store_inventory_presence(active_rows, rows, summary_checks, as_of, active_checks)
    yellow = PatternFill("solid", fgColor="FFFF00")
    for index in range(len(active_rows)):
        if index not in matched_active:
            final.cell(index + 2, DESC_COL).fill = yellow
    print(f"  Inventory SKUs with exact store evidence: {len(matched_active):,}", flush=True)
    print(f"  Yellow descriptions (no exact store evidence): {len(active_rows)-len(matched_active):,}", flush=True)
    first_store = len(warehouse_headers)+1
    for cells in final.iter_rows(min_row=2, min_col=first_store):
        for cell in cells:
            cell.number_format = '#,##0.00'
    apply_final_layout(final, warehouse_headers, headers)
    review_column = clean.max_column + 1
    clean.cell(1, review_column, "Review reason")
    clean.cell(1, review_column)._style = copy(clean.cell(1, 4)._style)
    clean.column_dimensions[get_column_letter(review_column)].width = 42
    review_count = 0
    for row_number, (summary_check, active_check) in enumerate(zip(summary_checks, active_checks), 2):
        if summary_check == 0:
            reason = "Not in summarized inventory"
        elif " | " in str(active_check):
            reason = "Ambiguous match"
        elif active_check == 0:
            reason = "Ambiguous summarized match" if " | " in str(summary_check) else "Inactive inventory"
        else:
            reason = ""
        clean.cell(row_number, review_column, reason)
        review_count += bool(reason)
    clean.auto_filter.ref = clean.dimensions
    print(f"  Stores clean rows: {len(rows):,}; rows with review reason: {review_count:,}", flush=True)
    for label, value in audit.items():
        print(f"  {label}: {value}", flush=True)
    return values


def apply_final_layout(sheet, warehouse_headers, stores_headers):
    """Two header rows matching layout.xlsx; months follow the rolling period."""
    sheet.insert_rows(2)
    for column in range(1,BASE_COLS+1):
        sheet.merge_cells(start_row=1, start_column=column, end_row=2, end_column=column)
    warehouse_end = len(warehouse_headers)
    stores_start = warehouse_end + 1
    stores_end = warehouse_end + len(stores_headers)
    for first,last,label,color in ((BASE_COLS+1,warehouse_end,"Warehouse Sales","548235"),
                                   (stores_start,stores_end,"Stores Sales","4472C4")):
        for column in range(first,last+1):
            old_header = warehouse_headers[column-1] if first == BASE_COLS+1 else stores_headers[column-stores_start]
            match = re.search(r"\b(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec) (\d{4})\b", old_header)
            label_month = calendar.month_name[datetime.strptime(match[1],"%b").month] if match else old_header
            sheet.cell(2,column,label_month)
            sheet.cell(1,column).value = None
        sheet.merge_cells(start_row=1,start_column=first,end_row=1,end_column=last)
        sheet.cell(1,first,label)
        sheet.cell(1,first).fill = PatternFill("solid",fgColor=Color(theme=3,tint=-0.249977111117893))
    for row in sheet.iter_rows(min_row=1,max_row=2):
        for cell in row:
            cell.font = Font(bold=True,color="FFFFFF")
            cell.alignment = Alignment(horizontal="center",vertical="center",wrap_text=True)
            if cell.row == 2 or cell.column <= BASE_COLS:
                cell.fill = PatternFill("solid",fgColor="17365D")
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 25
    sheet.auto_filter.ref = f"A2:{get_column_letter(stores_end)}{sheet.max_row}"


def format_output_tabs(target):
    # Style ids are set directly (fast) instead of assigning style objects to every cell.
    center_id = target._alignments.add(Alignment(horizontal="center", vertical="center"))
    wrap_id = target._alignments.add(Alignment(horizontal="center", vertical="center", wrap_text=True))
    for sheet in target:
        # Raw copies of the source exports stay exactly as they were. Only report tabs are centred.
        if sheet.title.endswith(" raw") or sheet.title.startswith("current "):
            continue
        alignments = target._alignments
        mapping = {}  # old alignment id -> centred id
        for cell in sheet._cells.values():
            style = cell._style
            if style is None:  # cell without any style yet: start from a default style array
                style = cell._style = StyleArray()
            old = style[5]  # alignmentId
            new = mapping.get(old)
            if new is None:
                current = alignments[old] if old < len(alignments) else None
                new = mapping[old] = wrap_id if (current is not None and current.wrap_text) else center_id
            style[5] = new
    final = target["final sheet"]
    edge = Side(style="thin", color="000000")
    border_id = target._borders.add(Border(left=edge, right=edge, top=edge, bottom=edge))
    for cell in final._cells.values():
        if cell._style is None:
            cell._style = StyleArray()
        cell._style[2] = border_id  # borderId
    for column in range(BASE_COLS+1, final.max_column + 1):
        final.column_dimensions[get_column_letter(column)].width = FINAL_SALES_COLUMN_WIDTH


def build_ordering_workbook(final, planning_cache, path):
    """Second workbook: tab 1 = copy of 'final sheet' with its formulas, tab 2 = ordering sheet.

    Ordering-sheet cells are formulas that point at tab 1, so editing Incoming QTY, Safety factor, Auto order?
    etc. on tab 1 updates the ordering sheet in Excel. Columns come from ORDERING_COLUMNS.
    """
    from openpyxl.worksheet.datavalidation import DataValidation
    headers = {}
    for column in range(1, final.max_column + 1):
        name = final.cell(2, column).value
        if name is None or isinstance(name, bool):
            name = final.cell(1, column).value
        if name is not None and str(name) not in headers:
            headers[str(name)] = column
    specs = []
    for item in ORDERING_COLUMNS:
        source, label = (item, item) if isinstance(item, str) else (item[0], item[1])
        if source in ("Adjustment", "Final order QTY"):
            specs.append((source, label, None))
            continue
        if source not in headers:
            raise ValueError(f"ORDERING_COLUMNS: '{source}' is not a column of the final sheet. Available: " + ", ".join(headers))
        specs.append((source, label, headers[source]))
    workbook = Workbook()
    workbook.remove(workbook.active)
    copy = workbook.create_sheet(final.title)
    copy_raw(final, copy)
    for validation in final.data_validations.dataValidation:
        new = DataValidation(type=validation.type, operator=validation.operator, formula1=validation.formula1,
                             formula2=validation.formula2, allow_blank=validation.allow_blank)
        new.errorStyle, new.showErrorMessage = validation.errorStyle, validation.showErrorMessage
        new.errorTitle, new.error = validation.errorTitle, validation.error
        for rng in validation.sqref.ranges:
            new.add(str(rng))
        copy.add_data_validation(new)
    orders = workbook.create_sheet(ORDERING_SHEET_NAME)
    last = final.max_row
    priority = headers.get("Priority")
    if ORDERING_ONLY_ORDER_ROWS and priority is None:
        raise ValueError("ORDERING_ONLY_ORDER_ROWS needs the 'Priority' column on the final sheet.")
    wanted = {str(x) for x in ORDERING_PRIORITIES}
    rows = []
    for r in range(3, last + 1):
        if ORDERING_ONLY_ORDER_ROWS and planning_cache.get(f"{get_column_letter(priority)}{r}") not in wanted:
            continue
        rows.append(r)
    header_fill = PatternFill("solid", fgColor="17365D")
    from openpyxl.comments import Comment as _Comment
    for position, (source, label, column) in enumerate(specs, 1):
        cell = orders.cell(1, position, label)
        if source == "Adjustment":
            cell.comment = _Comment("Type the COMPLETE quantity you want to order (e.g. suggested 10, you want 15: type 15). Blank = keep the suggestion, 0 = do not order.", "Planning")
        if source == "Final order QTY":
            cell.comment = _Comment("= Adjustment if you typed one, otherwise Suggested order QTY.", "Planning")
        cell.fill, cell.font = header_fill, Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        width = (final.column_dimensions[get_column_letter(column)].width or 16) if column else 16
        orders.column_dimensions[get_column_letter(position)].width = min(max(width, 12), 60)
    sheet_ref = "'" + final.title.replace("'", "''") + "'!"
    center = Alignment(horizontal="center", vertical="center")
    side = Side(style="thin", color="000000")
    edge_border = Border(left=side, right=side, top=side, bottom=side) if ORDERING_BORDERS else None
    if edge_border:
        for position in range(1, len(specs) + 1):
            orders.cell(1, position).border = edge_border
    where = {source: position for position, (source, label, column) in enumerate(specs, 1)}
    green, blue = PatternFill("solid", fgColor="E2F0D9"), Font(color="0000FF")
    suggested_pos, adjustment_pos = where.get("Suggested order QTY"), where.get("Adjustment")
    if "Final order QTY" in where and (suggested_pos is None or adjustment_pos is None):
        raise ValueError("ORDERING_COLUMNS: 'Final order QTY' needs both 'Suggested order QTY' and 'Adjustment' in the list.")
    for out_row, r in enumerate(rows, 2):
        for position, (source, label, column) in enumerate(specs, 1):
            if source == "Adjustment":
                cell = orders.cell(out_row, position)
                cell.fill, cell.font, cell.alignment, cell.number_format = green, blue, center, "0.##"
                if edge_border:
                    cell.border = edge_border
                continue
            if source == "Final order QTY":
                sug = f"{get_column_letter(suggested_pos)}{out_row}"
                adj = f"{get_column_letter(adjustment_pos)}{out_row}"
                cell = orders.cell(out_row, position, f'=IF(ISNUMBER({adj}),MAX(0,{adj}),IF(ISNUMBER({sug}),{sug},""))')
                cell.alignment, cell.number_format, cell.font = center, "0.##", Font(bold=True)
                if edge_border:
                    cell.border = edge_border
                continue
            ref = f"{sheet_ref}{get_column_letter(column)}{r}"
            cell = orders.cell(out_row, position, f'=IF({ref}="","",{ref})')
            cell.alignment = center
            if edge_border:
                cell.border = edge_border
            cell.number_format = final.cell(r, column).number_format
    orders.row_dimensions[1].height = 28
    if adjustment_pos and rows:
        from openpyxl.worksheet.datavalidation import DataValidation as _DV
        dv = _DV(type="decimal", operator="between", formula1="0", formula2="1000000", allow_blank=True)
        dv.errorStyle, dv.showErrorMessage, dv.errorTitle, dv.error = "stop", True, "Adjustment", "Enter the full quantity to order (0 or more), or leave blank to keep the suggestion."
        dv.add(f"{get_column_letter(adjustment_pos)}2:{get_column_letter(adjustment_pos)}{len(rows) + 1}")
        orders.add_data_validation(dv)
    orders.auto_filter.ref = f"A1:{get_column_letter(len(specs))}{max(1, len(rows) + 1)}"
    orders.sheet_view.showGridLines = False
    for ws in workbook:
        ws.freeze_panes = None
        ws.sheet_view.pane = None
        ws.sheet_view.selection = [Selection(activeCell="A1", sqref="A1")]
    workbook.calculation.fullCalcOnLoad = True
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.stem + ".tmp.xlsx")
    workbook.save(temporary)
    temporary.replace(path)
    workbook.close()
    print(f"  Ordering workbook: {path}  (tab 1 '{final.title}', tab 2 '{ORDERING_SHEET_NAME}' with {len(rows):,} rows x {len(specs)} columns)", flush=True)


def arrange_tabs(target):
    # Source tabs first, all clean tabs together, then inventory work and reports.
    for sheet in list(target):
        replacements = {"stores sales append raw": "current stores sales",
                        "sales append raw": "current WH Sales",
                        "purchase append raw": "current purchase"}
        for old, new in replacements.items():
            if sheet.title == old or sheet.title.startswith(old + " "):
                suffix = sheet.title[len(old):]
                sheet.title = (new + suffix)[:31]
                break
    source = [n for n in target.sheetnames if " raw" in n or n.startswith("current ")]
    clean = [n for n in target.sheetnames if n.endswith(" clean")]
    work = ["open PO", "master price book", "inv with PB desc", "inv summarized", "inv status lookup",
            "active inventory", "final sheet"]
    other = [n for n in target.sheetnames if n not in source + clean + work]
    ordered = source + clean + other + [n for n in work if n in target.sheetnames]
    for desired, name in enumerate(ordered):
        target.move_sheet(name, offset=desired-target.sheetnames.index(name))


def validate_demand_settings(cfg):
    """Fail early (before the long run) if a DEMAND setting is invalid."""
    weights = [float(w) for w in cfg["weights"]]
    if len(weights) != 3 or any(w < 0 for w in weights) or abs(sum(weights) - 1) > 1e-9:
        raise ValueError("DEMAND['weights'] must be 3 non-negative numbers that add up to 1 (newest window first).")
    if float(cfg["window_days"]) != int(cfg["window_days"]) or int(cfg["window_days"]) < 1:
        raise ValueError("DEMAND['window_days'] must be a whole number >= 1.")
    for key in ("lead_time", "default_incoming", "min_units_auto"):
        if float(cfg[key]) < 0:
            raise ValueError(f"DEMAND['{key}'] cannot be negative.")
    for key in ("reorder_below_days", "target_days", "default_order_multiple", "safety_factor", "new_item_factor", "open_po_qty_factor"):
        if float(cfg[key]) <= 0:
            raise ValueError(f"DEMAND['{key}'] must be greater than 0.")
    if not (0 < float(cfg["trend_falling"]) <= 1 <= float(cfg["trend_rising"])):
        raise ValueError("DEMAND trend limits need 0 < trend_falling <= 1 <= trend_rising.")
    a, b = cfg["abc_cutoffs"]
    if not (0 < float(a) < float(b) <= 1):
        raise ValueError("DEMAND['abc_cutoffs'] must look like (0.70, 0.90): 0 < A cut < B cut <= 1.")
    for key in ("class_reorder_below_days", "class_safety_factor"):
        for cls, value in cfg[key].items():
            if cls not in ("A", "B", "C") or float(value) <= 0:
                raise ValueError(f"DEMAND['{key}'] only accepts classes A/B/C with values > 0.")
    if int(cfg["new_item_store_periods"]) < 1:
        raise ValueError("DEMAND['new_item_store_periods'] must be >= 1.")


def open_po_rows(path):
    """Read the open PO workbook (first sheet that has a description and an open-quantity header).

    Accepted headers (any case/spacing): BOX / BOX UPC, IN / MASTER UPC / UPC, BRAND,
    DESCRIPTION / ITEM / PRODUCT/SERVICE, and the open quantity as OQ / OPEN QTY / BACKORDERED
    (or QTY minus RCVD). The file is read fresh on every run: it replaces the previous open PO.
    """
    aliases = {
        "box": ("box", "boxupc", "boxcode"),
        "master": ("in", "inner", "innerupc", "masterupc", "upc", "productupc"),
        "brand": ("brand",),
        "desc": ("description", "productservice", "productservicefullname", "item", "itemname", "product", "productname", "discription"),
        "open": ("oq", "openqty", "openquantity", "qtyopen", "backordered", "backorderedqty", "remainingqty", "qtyremaining", "openpoqty"),
        "ordered": ("qty", "quantity", "orderedqty", "qtyordered"),
        "received": ("rcvd", "received", "receivedqty", "qtyreceived", "rcvdqty"),
        "num": ("num", "ponumber", "po", "pono", "ponum"),
    }
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            header = None
            for row_number, cells in enumerate(sheet.iter_rows(min_row=1, max_row=20, values_only=True), 1):
                keys = {header_key(v): i for i, v in enumerate(cells) if v is not None}
                pick = {name: next((keys[a] for a in names if a in keys), None) for name, names in aliases.items()}
                if pick["desc"] is not None and (pick["open"] is not None or pick["ordered"] is not None):
                    header = (row_number, pick)
                    break
            if header is None:
                continue
            header_row, pick = header
            rows = []
            for row_number, cells in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
                get = lambda name: cells[pick[name]] if pick[name] is not None and pick[name] < len(cells) else None
                description = clean_description(trim(get("desc")))
                if not description or re.match(r"^(?:total\b|grand total\b)", description, re.I):
                    continue
                if pick["open"] is not None:
                    qty = number(get("open"), f"{sheet.title}!row {row_number}")
                else:
                    ordered = number(get("ordered"), f"{sheet.title}!row {row_number}")
                    received = number(get("received"), f"{sheet.title}!row {row_number}") if pick["received"] is not None else 0
                    qty = None if ordered is None else ordered - (received or 0)
                rows.append({"box": clean_upc(get("box")), "master": clean_upc(get("master")), "brand": trim(get("brand")),
                             "description": description, "qty": qty, "num": trim(get("num"))})
            return rows
        raise ValueError("Open PO file: no sheet with a DESCRIPTION header and an open quantity header (OQ / Open Qty / Backordered / Qty).")
    finally:
        workbook.close()


def purchase_costs(purchase_rows, checks):
    """Last and lowest buying price (with vendor) per active SKU from the cleaned Bill lines.

    `checks` is the lookup of every purchase line against active inventory (description, master/box UPC one
    code at a time, Price Book mapping). Lines that match more than one SKU, or none, are skipped.
    Rates of 0 or less are ignored. Last = latest bill date (later line wins a tie); lowest = smallest rate
    (the latest one wins a tie). Returns ({description_key: (last, last vendor, lowest, lowest vendor)}, stats).
    """
    best, stats = {}, Counter()
    for position, (row, check) in enumerate(zip(purchase_rows, checks)):
        if len(row) < 8 or not isinstance(row[7], (int, float)) or not math.isfinite(float(row[7])) or row[7] <= 0:
            stats["lines without a usable rate"] += 1
            continue
        if check == 0:
            stats["lines not matched to active inventory"] += 1
            continue
        if " | " in str(check):
            stats["lines matching several SKUs (skipped)"] += 1
            continue
        stamp = (row[0] or datetime.min, position)
        entry = best.setdefault(description_key(check), {"last": None, "low": None})
        if entry["last"] is None or stamp > entry["last"][0]:
            entry["last"] = (stamp, float(row[7]), row[6])
        if entry["low"] is None or (float(row[7]), -stamp[0].toordinal(), -position) < (entry["low"][1], -entry["low"][0][0].toordinal(), -entry["low"][0][1]):
            entry["low"] = (stamp, float(row[7]), row[6])
        stats["lines used"] += 1
    return {k: (v["last"][1], v["last"][2], v["low"][1], v["low"][2]) for k, v in best.items()}, stats


class InventoryMatcher:
    """Match a PO line to inventory rows with the same stages as the inventory/PB lookups.

    Stages (first stage with a hit wins): PO master -> inventory master, PO box -> inventory box,
    PO master -> inventory box, PO box -> inventory master, description -> description,
    description -> PB description, then the Price Book (PO UPC/description -> PB record ->
    inventory). Inventory cells holding several UPCs are checked one code at a time.
    """
    def __init__(self, rows, pb_indexes):
        self.rows, self.pb = rows, pb_indexes
        self.master, self.box, self.name, self.pbname = (defaultdict(set) for _ in range(4))
        for i, row in enumerate(rows):
            for code in upc_tokens(row[1]):
                self.master[code.casefold()].add(i)
            for code in upc_tokens(row[0]):
                self.box[code.casefold()].add(i)
            self.name[description_key(row[2])].add(i)
            for name in str(row[4] or "").split(" | "):
                if name and name != "-" and trim(name):
                    self.pbname[description_key(name)].add(i)

    def match(self, po):
        masters, boxes = upc_tokens(po["master"]), upc_tokens(po["box"])
        desc = description_key(po["description"])
        def hit(index, tokens):
            found = set()
            for code in tokens:
                found |= index.get(code.casefold(), set())
            return found
        stages = [("PO master to master", lambda: hit(self.master, masters)),
                  ("PO box to box", lambda: hit(self.box, boxes)),
                  ("PO master to box", lambda: hit(self.box, masters)),
                  ("PO box to master", lambda: hit(self.master, boxes)),
                  ("Description to description", lambda: set(self.name.get(desc, ()))),
                  ("Description to PB description", lambda: set(self.pbname.get(desc, ())))]
        for label, find in stages:
            found = find()
            if found:
                return found, label
        found = set()
        for field, tokens in (("master", masters), ("box", boxes)):
            for code in tokens:
                for record in self.pb[field].get(code.casefold(), []):
                    key = description_key(record["description"])
                    found |= self.pbname.get(key, set()) | self.name.get(key, set())
        if found:
            return found, "PO UPC to Price Book"
        for record in self.pb["description"].get(desc, []):
            for value in (record["master"], record["box"]):
                for code in upc_tokens(value):
                    found |= self.master.get(code.casefold(), set()) | self.box.get(code.casefold(), set())
        if found:
            return found, "PO description to Price Book"
        return set(), "Not found"

    def resolve(self, po):
        """Returns (index or None, method, candidates). Several candidates narrow to the exact description."""
        found, method = self.match(po)
        if len(found) > 1:
            exact = [i for i in found if description_key(self.rows[i][2]) == description_key(po["description"])]
            if len(exact) == 1:
                found = set(exact)
        if len(found) == 1:
            return next(iter(found)), method, found
        return None, method, found


def match_open_po(po_rows, summarized, active_rows, indexes, cfg=DEMAND):
    """Open PO lines -> active inventory first, then the rest of inventory (inactive SKUs that were re-ordered).

    Returns (tab rows, {description_key: units}, {summary index: PO description} for reactivated SKUs, stats).
    """
    factor = float(cfg["open_po_qty_factor"])
    active_matcher = InventoryMatcher(active_rows, indexes)
    active_keys = {description_key(a[2]) for a in active_rows}
    inactive_ix = [i for i, r in enumerate(summarized) if description_key(r[2]) not in active_keys]
    inactive_matcher = InventoryMatcher([summarized[i] for i in inactive_ix], indexes)
    incoming, reactivated, tab, stats = defaultdict(float), {}, [], Counter()
    for r in po_rows:
        qty = r["qty"]
        status, matched, method, counted = "", "", "", 0
        if qty is None or not math.isfinite(float(qty)) or qty <= 0:
            status = "No open quantity (ignored)"
        else:
            index, method, found = active_matcher.resolve(r)
            if index is not None:
                matched = active_rows[index][2]; status = "Counted - active SKU"
            elif found:
                matched = " | ".join(sorted(active_rows[i][2] for i in found)); status = "Ambiguous match (not counted)"
            else:
                index, method, found = inactive_matcher.resolve(r)
                if index is not None:
                    row = summarized[inactive_ix[index]]
                    matched = row[2]; status = "Counted - reactivated (was inactive)"
                    reactivated[inactive_ix[index]] = r["description"]
                elif found:
                    matched = " | ".join(sorted(inactive_matcher.rows[i][2] for i in found)); status = "Ambiguous match (not counted)"
                else:
                    status, method = "Not in inventory (not counted)", ""
            if status.startswith("Counted"):
                counted = qty * factor
                incoming[description_key(matched)] += counted
                stats["counted units"] += counted
        stats[status] += 1
        tab.append([r["box"], r["master"], r["brand"], r["description"], r["num"], qty, counted, matched, method, status])
    return tab, dict(incoming), reactivated, stats


def planning_inputs(path, cfg=DEMAND):
    """Optional workbook: description, Incoming QTY, Order multiple (boxes)."""
    if not path:
        return {}
    workbook = load_workbook(path, read_only=True, data_only=True)
    result = {}
    try:
        sheet = workbook.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        headers = {header_key(v): i for i,v in enumerate(next(iterator, ()))}
        if 'description' not in headers:
            raise ValueError('Planning inputs need a description header in row 1.')
        for row in iterator:
            def get(key):
                index = headers.get(key)
                return row[index] if index is not None and index < len(row) else None
            name = description_key(get('description'))
            if not name:
                continue
            incoming = number(get('incomingqty'), 'Planning incoming QTY')
            multiple = number(get('ordermultiple'), 'Planning order multiple')
            incoming = cfg["default_incoming"] if incoming is None else incoming
            multiple = cfg["default_order_multiple"] if multiple is None else multiple
            if not math.isfinite(float(incoming)) or incoming < 0 or not math.isfinite(float(multiple)) or multiple <= 0:
                raise ValueError('Planning incoming must be >=0 and order multiple must be >0: '+name)
            if name in result:
                raise ValueError('Duplicate planning input description: '+name)
            result[name] = (incoming, multiple)
        return result
    finally:
        workbook.close()


def add_demand_planning(sheet, active_rows, sales_rows, sales_checks, covered_months, as_of, overrides, cfg=DEMAND, store_values=None, incoming_notes=None, costs=None):
    """Rolling windows ending yesterday (3 windows of cfg['window_days']); quantity in warehouse units.

    Every SKU gets an ABC class (volume) and an Item class (behaviour). Only items whose
    behaviour is reliable (Regular, or New with store sales) are ordered automatically;
    the rest show a Reference order QTY plus a Review priority.
    Missing history/ambiguous WH allocation blocks automated order calculation.
    In-stock history is unavailable: stockout-suppressed sales need manual review.
    Returns formula caches for display before Excel first recalculates.
    """
    from openpyxl.comments import Comment
    from openpyxl.worksheet.datavalidation import DataValidation
    validate_demand_settings(cfg)
    W = int(cfg['window_days'])
    lead0, base_trigger, target0 = cfg['lead_time'], cfg['reorder_below_days'], cfg['target_days']
    w_dec = [Decimal(str(x)) for x in cfg['weights']]
    w_txt = [repr(float(x)) for x in cfg['weights']]
    nf = float(cfg['new_item_factor']); nf_txt = repr(nf)
    rising, falling = Decimal(str(cfg['trend_rising'])), Decimal(str(cfg['trend_falling']))
    rising_txt, falling_txt = repr(float(cfg['trend_rising'])), repr(float(cfg['trend_falling']))
    end = as_of - timedelta(days=1)
    windows = [(end-timedelta(days=W-1+W*i), end-timedelta(days=W*i)) for i in range(3)]
    coverage = [all((day.year,day.month) in covered_months for day in
                    (start+timedelta(days=j) for j in range(W))) for start,finish in windows]
    by_name = {description_key(r[2]):i for i,r in enumerate(active_rows)}
    qty = defaultdict(Decimal)
    mtd = defaultdict(Decimal)
    issues = defaultdict(set)
    latest_date = max((r[1].date() for r in sales_rows if r[1] and r[1].date() <= end), default=None)
    for row,check in zip(sales_rows,sales_checks):
        if row[1] is None:
            continue
        day = row[1].date()
        if not windows[2][0] <= day <= end:
            continue
        candidates = {by_name[key] for key in (description_key(n) for n in str(check).split(' | ')) if key in by_name} if check != 0 else set()
        if len(candidates) > 1:
            for i in candidates:
                issues[i].add('Ambiguous warehouse sales')
            continue
        if not candidates:
            continue
        i = next(iter(candidates))
        if row[6] is None or not math.isfinite(float(row[6])):
            issues[i].add('Missing warehouse quantity')
            continue
        window = (end-day).days//W
        qty[i,window] += Decimal(str(row[6]))
        if (day.year,day.month) == (end.year,end.month):
            mtd[i] += Decimal(str(row[6]))

    # ---------- classification (Python values; Auto order? stays editable in Excel) ----------
    recents = []
    for i in range(len(active_rows)):
        recents.append([max(Decimal(0),qty[i,j]) if coverage[j] and not issues[i] else None for j in range(3)])
    totals = [sum(r) if all(v is not None for v in r) else None for r in recents]
    grand = sum(t for t in totals if t is not None)
    abc = ['-'] * len(active_rows)
    running = Decimal(0)
    cut_a, cut_b = Decimal(str(cfg['abc_cutoffs'][0])), Decimal(str(cfg['abc_cutoffs'][1]))
    for i in sorted((i for i,t in enumerate(totals) if t is not None), key=lambda i: (-totals[i], i)):
        running += totals[i]
        share = running/grand if grand > 0 else Decimal(1)
        abc[i] = 'A' if (totals[i] > 0 and share <= cut_a) else 'B' if (totals[i] > 0 and share <= cut_b) else 'C'
    periods = int(cfg['new_item_store_periods'])
    def stores_selling(i):
        if not store_values:
            return False
        vals = [v for v in store_values[i][-periods:] if isinstance(v,(int,float))]
        return any(v > 0 for v in vals)
    item_class, why_not = [], []
    min_units = Decimal(str(cfg['min_units_auto']))
    for i in range(len(active_rows)):
        q = recents[i]
        if any(v is None for v in q):
            item_class.append('Review data'); why_not.append('')
        elif totals[i] == 0:
            item_class.append('No demand'); why_not.append('')
        elif q[0] > 0 and q[1] > 0:
            if totals[i] >= min_units:
                item_class.append('Regular'); why_not.append('')
            else:
                item_class.append('Slow regular'); why_not.append(f'Slow regular: under {cfg["min_units_auto"]} units in {3*W} days')
        elif q[0] > 0 and q[1] == 0 and q[2] == 0:
            if stores_selling(i):
                item_class.append('New - stores selling'); why_not.append('')
            else:
                item_class.append('New - no store sales'); why_not.append('New/one-off item: sold only in the last window and stores show no sale')
        elif q[0] > 0:
            item_class.append('Intermittent')
            why_not.append('' if abc[i] in ('A','B') else 'Intermittent: skipped a recent window, not auto-ordered')
        elif ((q[1] > 0) != (q[2] > 0)) or (abc[i] != 'A' and float(max(q[1], q[2])) >= 0.8 * float(q[1] + q[2] + q[0])):
            item_class.append('One-off (stopped)')
            why_not.append('One-off item: sold in a single window and not since, not auto-ordered')
        else:
            item_class.append('Lapsed')
            why_not.append('Lapsed fast item (A): no sale in the latest window, check if it stocked out, not auto-ordered' if abc[i] == 'A'
                           else 'Lapsed: no sale in the latest window, not auto-ordered')
    AUTO_CLASSES = ('Regular', 'New - stores selling')
    def is_auto(i):
        return item_class[i] in AUTO_CLASSES or (item_class[i] == 'Intermittent' and abc[i] in ('A','B'))

    first = sheet.max_column + 1
    labels = ['WH window 1','WH window 2','WH window 3','Daily demand','Current month avg/day',
              'WH trend','Lead days','Reorder below days','Target cover days',
              'Incoming QTY','Order multiple','Days cover','Stock at arrival',
              'Suggested order QTY','Priority','Planning notes',
              'ABC class','Item class','Auto order?','Safety factor','Reference order QTY',
              'Last buying price','Last vendor','Lowest buying price','Lowest price vendor']
    shown = {'WH window 1': f'WH last {W} days', 'WH window 2': f'WH previous {W} days',
             'WH window 3': f'WH days {2*W+1}-{3*W}'}
    columns = {name:first+i for i,name in enumerate(labels)}
    for i,label in enumerate(labels):
        cell = sheet.cell(2,first+i,shown.get(label,label))
        cell.fill = PatternFill('solid',fgColor='17365D')
        cell.font = Font(bold=True,color='FFFFFF')
        cell.alignment = Alignment(horizontal='center',vertical='center',wrap_text=True)
    sheet.merge_cells(start_row=1,start_column=first,end_row=1,end_column=first+len(labels)-1)
    sheet.cell(1,first,'Demand Planning')
    sheet.cell(1,first).fill = PatternFill('solid',fgColor='375623')
    sheet.cell(1,first).font = Font(bold=True,color='FFFFFF')
    for i,(start,finish) in enumerate(windows):
        sheet.cell(2,first+i).comment = Comment(f'{start} through {finish}, {W} calendar days. Net invoice quantity; negative net demand is clipped to zero.','Planning')
    sheet.cell(2,columns['Daily demand']).comment = Comment(f'Weights {w_txt[0]} / {w_txt[1]} / {w_txt[2]} on the three {W}-day windows (newest first); divide by {W}. For "New - stores selling" items: {nf_txt} x units sold in the latest window / {W} (expected store re-orders). Requires complete sales exports. Stockout days cannot be adjusted without stock history.','Planning')
    sheet.cell(2,columns['Current month avg/day']).comment = Comment(f'Units sold from {end.replace(day=1)} through {end} divided by {end.day} days (month to date, yesterday is the last full day). For comparison with Daily demand: a higher value means this month is running above the weighted average. Not used in the order calculation.','Planning')
    sheet.cell(2,columns['Last buying price']).comment = Comment('Rate of the latest Bill line for this SKU in the purchase exports (6 months + current month). Vendor = the vendor of that Bill. Blank = no priced purchase found.','Planning')
    sheet.cell(2,columns['Lowest buying price']).comment = Comment('Smallest Rate on any Bill line for this SKU in the purchase exports, with the vendor who gave it.','Planning')
    sheet.cell(2,columns['Incoming QTY']).comment = Comment(f'Editable. Default {cfg["default_incoming"]} assumes no pending receipts. Count only confirmed incoming in warehouse units arriving by the {lead0}-day receipt date. Use planning_inputs workbook to persist per-SKU entries across script runs.','Planning')
    sheet.cell(2,columns['Target cover days']).comment = Comment(f'{target0} days stock AFTER delivery. Order target includes consumption during the {lead0}-day lead time.','Planning')
    sheet.cell(2,columns['Reorder below days']).comment = Comment(f'Default by ABC class: A={cfg["class_reorder_below_days"].get("A",base_trigger)}, B={cfg["class_reorder_below_days"].get("B",base_trigger)}, C={cfg["class_reorder_below_days"].get("C",base_trigger)} days. Faster movers are reordered earlier.','Planning')
    sheet.cell(2,columns['ABC class']).comment = Comment(f'By units sold in the {3*W} days: A = top SKUs making up {cfg["abc_cutoffs"][0]:.0%} of volume, B = next up to {cfg["abc_cutoffs"][1]:.0%}, C = rest.','Planning')
    sheet.cell(2,columns['Item class']).comment = Comment('Regular: sold in the last two windows. New - stores selling / no store sales: sold only in the latest window (new or one-off item); stores decide. Intermittent: skipped a recent window (auto only if A/B). One-off (stopped): sold in one earlier window only. Lapsed: sold earlier in two windows but not in the latest. Slow regular: very low volume.','Planning')
    sheet.cell(2,columns['Auto order?']).comment = Comment('Yes = Suggested order QTY is calculated. No = Suggested order QTY is 0 and the Reference order QTY shows what would be ordered. Editable: change to Yes/No per SKU.','Planning')
    sheet.cell(2,columns['Safety factor']).comment = Comment(f'Multiplies the demand over (lead + target) days. Default by class: A={cfg["class_safety_factor"].get("A",cfg["safety_factor"])}, B={cfg["class_safety_factor"].get("B",cfg["safety_factor"])}, C={cfg["class_safety_factor"].get("C",cfg["safety_factor"])}. Higher = fewer stock-outs, more stock.','Planning')
    for label in ['Lead days','Reorder below days','Target cover days','Incoming QTY','Order multiple','Safety factor']:
        c = get_column_letter(columns[label]); minimum = 0 if label in ['Lead days','Incoming QTY'] else 0.000001
        dv = DataValidation(type='decimal',operator='greaterThanOrEqual',formula1=str(minimum),allow_blank=False)
        dv.errorTitle='Invalid planning input';dv.error='Enter a valid non-negative quantity/day value (order multiple, safety factor must be positive).';dv.showErrorMessage=True;dv.errorStyle='stop'
        sheet.add_data_validation(dv);dv.add(f'{c}3:{c}{len(active_rows)+2}')
    c = get_column_letter(columns['Auto order?'])
    dv = DataValidation(type='list',formula1='"Yes,No"',allow_blank=False); dv.errorStyle='stop'; dv.showErrorMessage=True
    sheet.add_data_validation(dv); dv.add(f'{c}3:{c}{len(active_rows)+2}')
    caches = {}
    counts, class_counts = Counter(), Counter()
    for i,active in enumerate(active_rows):
        r=i+3
        c={label:f'{get_column_letter(col)}{r}' for label,col in columns.items()}
        recent=recents[i]
        for j,value in enumerate(recent):
            sheet.cell(r,first+j,float(value) if value is not None else 'Review data')
        incoming,multiple=overrides.get(description_key(active[2]),(cfg['default_incoming'],cfg['default_order_multiple']))
        trigger0=cfg['class_reorder_below_days'].get(abc[i],base_trigger)
        safety0=cfg['class_safety_factor'].get(abc[i],cfg['safety_factor'])
        auto_value='Yes' if is_auto(i) else 'No'
        inputs={'Lead days':lead0,'Reorder below days':trigger0,'Target cover days':target0,'Incoming QTY':incoming,'Order multiple':multiple,'Safety factor':safety0}
        for label,value in inputs.items():
            cell=sheet.cell(r,columns[label],value);cell.fill=PatternFill('solid',fgColor='E2F0D9');cell.font=Font(color='0000FF')
        sheet.cell(r,columns['ABC class'],abc[i])
        sheet.cell(r,columns['Item class'],item_class[i])
        cell=sheet.cell(r,columns['Auto order?'],auto_value);cell.fill=PatternFill('solid',fgColor='E2F0D9');cell.font=Font(color='0000FF')
        q0,q1,q2=[c[x] for x in labels[:3]]
        daily=c['Daily demand'];lead=c['Lead days'];trigger=c['Reorder below days'];target=c['Target cover days'];inc=c['Incoming QTY'];mult=c['Order multiple'];cover=c['Days cover'];order=c['Suggested order QTY'];stock=f'{QTY_COL_LETTER}{r}'
        cls=c['Item class'];auto=c['Auto order?'];safety=c['Safety factor'];ref=c['Reference order QTY']
        def formula(label,expression,value):
            cell=sheet.cell(r,columns[label],expression);caches[cell.coordinate]=value
        if all(v is not None for v in recent):
            if item_class[i]=='New - stores selling':
                d=nf*float(recent[0])/W
            else:
                d=float(sum(v*w for v,w in zip(recent,w_dec)))/W
        else:
            d=None
        formula('Daily demand',f'=IF(COUNT({q0}:{q2})=3,IF({cls}="New - stores selling",{nf_txt}*{q0}/{W},({q0}*{w_txt[0]}+{q1}*{w_txt[1]}+{q2}*{w_txt[2]})/{W}),"")',d if d is not None else '')
        tr='Review data' if any(v is None for v in recent[:2]) else ('No movement' if recent[0]==0 else 'New / resumed') if recent[1]==0 else 'Rising' if recent[0]/recent[1]>=rising else 'Falling' if recent[0]/recent[1]<=falling else 'Stable'
        month_ok = (end.year,end.month) in covered_months and not issues[i]
        cm = max(Decimal(0),mtd[i])/Decimal(end.day) if month_ok else None
        cell = sheet.cell(r,columns['Current month avg/day'],float(cm) if cm is not None else 'Review data')
        formula('WH trend',f'=IF(COUNT({q0}:{q1})<2,"Review data",IF({q1}=0,IF({q0}=0,"No movement","New / resumed"),IF({q0}/{q1}>={rising_txt},"Rising",IF({q0}/{q1}<={falling_txt},"Falling","Stable"))))',tr)
        stock_value=active[3]
        if isinstance(stock_value,(int,float)) and math.isfinite(float(stock_value)) and stock_value<0:
            stock_value=0  # negative stock is shown as 0 in the final sheet, plan with 0
        valid_stock=isinstance(stock_value,(int,float)) and math.isfinite(float(stock_value)) and stock_value>=0
        days=stock_value/d if valid_stock and d is not None and d>0 else None
        formula('Days cover',f'=IF(OR(NOT(ISNUMBER({daily})),{daily}<=0,NOT(ISNUMBER({stock})),{stock}<0),"",{stock}/{daily})',days if days is not None else '')
        at_arrival=max(0,stock_value-lead0*d+incoming) if days is not None else ''
        formula('Stock at arrival',f'=IF(NOT(ISNUMBER({cover})),"",MAX(0,{stock}-{daily}*{lead}+{inc}))',at_arrival)
        reference=math.ceil(round(max(0,d*(target0+lead0)*safety0-stock_value-incoming)/multiple,8))*multiple if days is not None and days<trigger0 else 0 if days is not None else ''
        formula('Reference order QTY',f'=IF(OR(NOT(ISNUMBER({cover})),COUNT({lead}:{mult})<>5,NOT(ISNUMBER({safety})),{lead}<0,{trigger}<=0,{target}<=0,{inc}<0,{mult}<=0,{safety}<=0),"",IF({cover}>={trigger},0,ROUNDUP(ROUND(MAX(0,{daily}*({target}+{lead})*{safety}-{stock}-{inc})/{mult},8),0)*{mult}))',reference)
        suggested=(reference if auto_value=='Yes' else 0) if reference!='' else ''
        formula('Suggested order QTY',f'=IF(NOT(ISNUMBER({ref})),"",IF({auto}="Yes",{ref},0))',suggested)
        if d is None or not valid_stock: priority='Review data'
        elif d<=0: priority='No demand - review'
        elif auto_value!='Yes': priority='Review - not auto' if days<trigger0 else 'Hold'
        else: priority='Urgent - order' if days<lead0 and suggested>0 else 'Urgent - expedite incoming' if days<lead0 else 'Order' if suggested>0 else 'Incoming covers target' if days<trigger0 else 'Hold'
        formula('Priority',f'=IF(NOT(ISNUMBER({daily})),"Review data",IF(OR(NOT(ISNUMBER({stock})),{stock}<0),"Review data",IF({daily}<=0,"No demand - review",IF(NOT(ISNUMBER({ref})),"Review inputs",IF({auto}<>"Yes",IF({cover}<{trigger},"Review - not auto","Hold"),IF({cover}<{lead},IF({order}>0,"Urgent - order","Urgent - expedite incoming"),IF({order}>0,"Order",IF({cover}<{trigger},"Incoming covers target","Hold"))))))))',priority)
        notes=sorted(issues[i])
        if incoming_notes and description_key(active[2]) in incoming_notes:notes.append(incoming_notes[description_key(active[2])])
        if why_not[i]:notes.append(why_not[i])
        if item_class[i]=='New - stores selling':notes.append(f'New item: expected demand = {nf_txt} x launch units, stores are selling')
        if not all(coverage):notes.append(f'Incomplete {3*W}-day WH history')
        if not valid_stock:notes.append('Review warehouse stock')
        if any(qty[i,j]<0 for j in range(3)):notes.append('Negative net sales clipped to zero')
        if d==0:notes.append('No recent demand: check new item or stockout')
        if cfg['show_export_note'] and (latest_date is None or latest_date<end):notes.append(f'Latest WH sale: {latest_date}; confirm export through {end}')
        sheet.cell(r,columns['Planning notes'],'; '.join(notes))
        cost=(costs or {}).get(description_key(active[2]))
        for label,value in zip(['Last buying price','Last vendor','Lowest buying price','Lowest price vendor'],cost or ('','','','')):
            cell=sheet.cell(r,columns[label],value)
            if 'price' in label:
                cell.number_format='#,##0.00'
        for label in ['Daily demand','Current month avg/day','Days cover','Stock at arrival','Suggested order QTY','Reference order QTY']:
            sheet.cell(r,columns[label]).number_format='0.00'
        counts[priority]+=1
        class_counts[item_class[i]]+=1
    sheet.auto_filter.ref=f'A2:{get_column_letter(sheet.max_column)}{sheet.max_row}'
    sheet.row_dimensions[2].height=42
    print(f'  Planning: last completed day {end}; {3*W} days in 3 windows of {W}; weights {"/".join(w_txt)}; lead {lead0}; target {target0} after receipt',flush=True)
    print('  Item classes: '+str(dict(class_counts)),flush=True)
    print('  Planning priorities: '+str(dict(counts)),flush=True)
    return caches


def run(paths):
    validate_demand_settings(DEMAND)
    source_paths = {k: Path(paths[k]).expanduser().resolve() for k in ("inventory", "purchase", "sales", "pricebook")}
    output = Path(paths["output"]).expanduser().resolve()
    history_path = Path(paths.get("history") or str(output) + ".history.json").expanduser().resolve()
    if history_path == output or history_path in source_paths.values():
        raise ValueError("History path must differ from the workbook and inputs.")
    history = {"version": 1}
    if history_path.exists():
        history = json.loads(history_path.read_text(encoding="utf-8"))
        if history.get("version") != 1:
            raise ValueError("Unsupported append history version.")
    # Remove previous fuzzy cache; persisted purchase/sales append data is retained.
    history.pop("stores_fuzzy_aliases", None)
    for kind, path in source_paths.items():
        if not path.is_file():
            raise FileNotFoundError(f"{kind} file not found: {path}")
        if path == output:
            raise ValueError("Output must not overwrite any input file.")
    target = Workbook()
    target.remove(target.active)
    print("Loading Master Price Book...", flush=True)
    pb_records = read_pricebook(source_paths["pricebook"])
    indexes = pricebook_indexes(pb_records)
    print(f"  Price Book source records: {len(pb_records):,}", flush=True)
    inventory_clean_rows = []
    purchase_clean_rows = []
    sales_clean_rows = []
    sales_base_sheet = None
    for kind in ("inventory", "purchase", "sales"):
        print(f"\nReading {kind}: {source_paths[kind]}", flush=True)
        source = load_workbook(source_paths[kind], data_only=False)
        cached = load_workbook(source_paths[kind], data_only=True, read_only=True)
        try:
            sheet, header_row, cols = find_source(cached, kind)
            raw = target.create_sheet(kind + " raw")
            print(f"  Preserving original sheet: {sheet.title}", flush=True)
            copy_raw(source[sheet.title], raw)
            clean = target.create_sheet(kind + " clean")
            stats = Counter()
            headers, rows = (inventory_rows(sheet, header_row, cols, stats) if kind == "inventory"
                             else transaction_rows(sheet, header_row, cols, kind, stats))
            if kind in ("purchase", "sales"):
                append_name = paths.get(kind + "_append", "")
                append_rows, coverage = [], None
                if append_name:
                    append_path = Path(append_name).expanduser().resolve()
                    if append_path == output or append_path == history_path or append_path in source_paths.values():
                        raise ValueError("Append paths must differ from initial files/output/history.")
                    append_source = load_workbook(append_path, data_only=False)
                    append_cached = load_workbook(append_path, read_only=True, data_only=True)
                    try:
                        a_sheet, a_header, a_cols = find_source(append_cached, kind)
                        append_headers, append_rows = transaction_rows(a_sheet, a_header, a_cols, kind, Counter())
                        coverage = append_date_range(a_sheet, kind, a_cols, append_rows)
                        copy_raw(append_source[a_sheet.title], target.create_sheet(kind + " append raw"))
                        print(f"  Append replaces dates {coverage[0]} through {coverage[1]} ({len(append_rows):,} qualifying rows)", flush=True)
                    finally:
                        append_source.close()
                        append_cached.close()
                rows = merge_daily_history(rows, append_rows, kind, history, coverage)
            write_clean(clean, headers, rows)
            if kind == "purchase":
                purchase_clean_rows = rows
            if kind == "sales":
                sales_clean_rows = rows
                # This copied worksheet survives closing the source workbook.
                sales_base_sheet = raw
            if kind == "inventory":
                inventory_clean_rows = rows
                # Store real Excel PROPER formulas; cached results added on save.
                for row_number, row in enumerate(rows, 2):
                    clean.cell(row_number, 3, '=PROPER("' + row[2].replace('"', '""') + '")')
            print(f"  Clean rows: {len(rows):,}", flush=True)
            qty_index = headers.index("QTY")
            print(f"  Clean quantity total: {sum(r[qty_index] or 0 for r in rows):,.4f}", flush=True)
            for label, count in stats.items():
                print(f"  {label}: {count:,}", flush=True)
            if kind == "inventory":
                duplicate_rows = len(rows) - len({r[2].casefold() for r in rows})
                print(f"  Repeated cleaned descriptions: {duplicate_rows:,} (rows retained)", flush=True)
        finally:
            source.close()
            cached.close()
    proper_inventory = [[r[0], r[1], proper(r[2]), r[3]] for r in inventory_clean_rows]
    print("\nMapping inventory to Master Price Book...", flush=True)
    mapped, lookup_counts = lookup_inventory(proper_inventory, indexes)
    mapped_sheet = target.create_sheet("inv with PB desc")
    write_clean(mapped_sheet,
                ["box upc", "master upc", "description", "QTY", "PB description", "Department", "Brand", "Pack of",
                 "Lookup method", "Matched upc", "Lookup status"],
                [r[:5] + [r[8], r[9], r[10]] + r[5:8] for r in mapped])
    for label, count in lookup_counts.items():
        print(f"  {label}: {count:,}", flush=True)
    pb_sheet = target.create_sheet("master price book")
    write_clean(pb_sheet, ["box upc", "master upc", "PB description", "Brand", "Department", "Pack of", "Source sheet", "Source row", "Source block", "Lookup priority"],
                [[r["box"], r["master"], r["description"], r["brand"], r["department"], r["pack"], r["sheet"], r["row"], r["block"], r["priority"]] for r in pb_records])
    print("\nSummarizing inventory by cleaned QB description...", flush=True)
    summarized = summarize_inventory(mapped)
    summary_sheet = target.create_sheet("inv summarized")
    write_clean(summary_sheet,
                ["box upc", "master upc", "description", "QTY", "PB description"],
                [row[:5] for row in summarized])
    print(f"  Summarized rows: {len(summarized):,}", flush=True)
    print("  Inventory quantity total reconciled.", flush=True)
    print("\nChecking summarized inventory against cleaned purchases...", flush=True)
    purchase_checks, reverse_matches = reverse_purchase_check(summarized, purchase_clean_rows, indexes)
    print(f"  Purchase rows matched back to inventory: {sum(value != 0 for value in purchase_checks):,}", flush=True)
    print(f"  Purchase rows with lookup check = 0: {purchase_checks.count(0):,}", flush=True)
    audit_rows, active_rows = purchase_activity(summarized, purchase_clean_rows, reverse_matches)
    # Open PO: read fresh each run (a new file fully replaces the old one, nothing is stored in history).
    overrides = planning_inputs(paths.get("planning_inputs", ""), DEMAND)
    incoming_notes, po_incoming = {}, {}
    po_path = paths.get("open_po", "")
    if po_path:
        po_path = Path(po_path).expanduser().resolve()
        if not po_path.is_file():
            raise FileNotFoundError(f"open PO file not found: {po_path}")
        if po_path == output or po_path in source_paths.values():
            raise ValueError("Open PO path must differ from the workbook and the other inputs.")
        print("\nReading open PO and matching it to inventory...", flush=True)
        po_rows = open_po_rows(po_path)
        po_tab, po_incoming, reactivated, po_stats = match_open_po(po_rows, summarized, active_rows, indexes, DEMAND)
        write_clean(target.create_sheet("open PO"),
                    ["box upc", "master upc", "Brand", "description", "PO num", "Open QTY", "Counted QTY (x factor)",
                     "Matched inventory description", "Lookup method", "Match status"], po_tab)
        if reactivated:
            # Inactive SKUs that were re-ordered join the active list (and so the final sheet) before sales/store mapping.
            audit_rows, active_rows = purchase_activity(summarized, purchase_clean_rows, reverse_matches, force_active=reactivated)
        print(f"  Open PO lines: {len(po_rows):,}", flush=True)
        for label, value in po_stats.items():
            print(f"  {label}: {value:,.0f}", flush=True)
    else:
        print("\nOpen PO: no file set, Incoming QTY stays at the default.", flush=True)
    base_headers = ["box upc", "master upc", "description", "QTY", "PB description"]
    write_clean(target.create_sheet("inv status lookup"),
                base_headers + ["Status", "Lookup method", "Matched UPC or description"], audit_rows)
    write_clean(target.create_sheet("active inventory"), base_headers + ["Status", "Department", "Brand", "Pack of"], active_rows)
    purchase_active_checks, _ = reverse_purchase_check(active_rows, purchase_clean_rows, indexes)
    costs, cost_stats = purchase_costs(purchase_clean_rows, purchase_active_checks)
    print(f"  Buying prices found for {len(costs):,} active SKUs", flush=True)
    for label, value in cost_stats.items():
        print(f"  Purchase price lookup, {label}: {value:,}", flush=True)
    add_lookup_columns(target["purchase clean"], purchase_checks, purchase_active_checks)
    # Sales rows have an extra Store name column. Pass the six item/date fields
    # to the same exact matching engine; all original sales rows are retained.
    sales_items = [row[1:7] for row in sales_clean_rows]
    sales_summary_checks, _ = reverse_purchase_check(summarized, sales_items, indexes)
    sales_active_checks, _ = reverse_purchase_check(active_rows, sales_items, indexes)
    add_lookup_columns(target["sales clean"], sales_summary_checks, sales_active_checks)
    as_of = date.fromisoformat(AS_OF_DATE) if AS_OF_DATE else datetime.now(timezone(timedelta(hours=5))).date()
    covered_months = sales_coverage_months(sales_base_sheet, sales_clean_rows, history)
    sales_headers, monthly_rows, sales_audit = monthly_sales_rows(active_rows, sales_clean_rows, sales_active_checks, covered_months, as_of)
    print(f"  Monthly sales report as of {as_of}", flush=True)
    for label, value in sales_audit.items():
        print(f"  {label}: {value}", flush=True)
    print(f"  Sales rows not found in summarized: {sales_summary_checks.count(0):,}", flush=True)
    print(f"  Sales rows not found in active: {sales_active_checks.count(0):,}", flush=True)
    dates = [r[0] for r in purchase_clean_rows if r[0] is not None]
    if dates:
        print(f"  Purchase Bill dates: {min(dates):%m/%d/%Y} through {max(dates):%m/%d/%Y}", flush=True)
    print(f"  Active: {len(active_rows):,}", flush=True)
    print(f"  Inactive: {len(audit_rows) - len(active_rows):,}", flush=True)
    print("\nPreparing store sales and cross checks...", flush=True)
    store_values = add_stores_reports(paths, target, summarized, active_rows, indexes, history, as_of, sales_headers, monthly_rows)
    for key, qty in po_incoming.items():
        inc, mult = overrides.get(key, (DEMAND["default_incoming"], DEMAND["default_order_multiple"]))
        overrides[key] = (inc + qty, mult)
        incoming_notes[key] = f"Open PO: {qty:g} units incoming"
    planning_cache = add_demand_planning(target["final sheet"], active_rows, sales_clean_rows, sales_active_checks, covered_months, as_of,
                                         overrides, DEMAND, store_values, incoming_notes, costs)
    arrange_tabs(target)
    format_output_tabs(target)
    if not INCLUDE_RAW_TABS:
        for name in list(target.sheetnames):
            if name.endswith(" raw") or name.startswith("current "):
                del target[name]
    for ws in target:
        # No frozen panes anywhere. Resetting only freeze_panes leaves a stale <selection pane="bottomLeft">
        # behind, which makes Excel report a corrupted file, so the selection is replaced as well.
        ws.freeze_panes = None
        ws.sheet_view.pane = None
        ws.sheet_view.selection = [Selection(activeCell="A1", sqref="A1")]
    ordering_name = paths.get("output_ordering", "")
    if ordering_name:
        ordering_path = Path(ordering_name).expanduser().resolve()
        if ordering_path == output or ordering_path in source_paths.values():
            raise ValueError("Ordering workbook path must differ from the main output and the inputs.")
        print("\nBuilding ordering workbook...", flush=True)
        build_ordering_workbook(target["final sheet"], planning_cache, ordering_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    # Save atomically so a failed save does not leave a half-written deliverable.
    temporary = output.with_name(output.stem + ".tmp.xlsx")
    print(f"\nSaving: {output}", flush=True)
    # Formulas are saved WITHOUT patching the XML afterwards (that patching made Excel ask to repair the
    # file and doubled the save time). Excel calculates every formula when the file is opened.
    target.calculation.fullCalcOnLoad = True
    target.save(temporary)
    temporary.replace(output)
    history_path.parent.mkdir(parents=True, exist_ok=True)
    history_tmp = history_path.with_name(history_path.name + ".tmp")
    history_tmp.write_text(json.dumps(history, ensure_ascii=False), encoding="utf-8")
    history_tmp.replace(history_path)
    target.close()
    print(f"[OK] {len(target.sheetnames)} tabs saved. Append history: {history_path}", flush=True)
    return output


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    for key, default in CONFIG.items():
        parser.add_argument("--" + key, default=default)
    args = vars(parser.parse_args())
    try:
        run(args)
    except Exception as exc:
        print(f"\n[ERROR] {exc}", file=sys.stderr, flush=True)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())