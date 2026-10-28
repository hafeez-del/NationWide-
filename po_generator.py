import argparse
import re
import sys
from collections import OrderedDict,defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ============================== SETTINGS (change these) ==============================
CONFIG = {
    "ordering_workbook": r"E:\Downloads\Automate\Warehouse_Ordering_Sheet_28.xlsx",
    "output_folder": r"E:\Downloads\Automate\PO_new2",
}

PO_PRICE_BASIS = "last"          # which VENDOR to use (no prices are printed any more):
                                 # "last"  = Last vendor
                                 # "lowest" = Lowest price vendor
ORDERING_SHEET = "ordering sheet"  # tab of the ordering workbook to read
PO_COLUMN = "PO"                 # header of the column in which you write PO
PO_MARK = "PO"                   # the text that selects a row (not case sensitive)

# Headers of the ordering sheet that the script reads. Change them if you renamed a column there.
COLUMNS = {
    "description": "description",
    "qty": "Final order QTY",
    "last_price": "Last buying price",
    "last_vendor": "Last vendor",
    "lowest_price": "Lowest buying price",
    "lowest_vendor": "Lowest price vendor",
    "box_upc": "box upc",        # optional columns, only shown when SHOW_UPC_* is True
    "master_upc": "master upc",
    "vendor_override": "PO vendor",   # optional: type a NEW vendor here -> used instead of Last/Lowest vendor
}

# Company details printed at the top of every PO PDF. Empty lines are skipped.
COMPANY = {
    "name": "Nation Wide Whole Sale LLC",
    "lines": ["", "", ""],      # e.g. ["Street address", "City, State ZIP", "Phone / Email"]
}
PO_PREFIX = "PO"                 # PO number = PREFIX-YYYYMMDD-NN (NN follows the vendor names A to Z)
PO_DATE = ""                     # YYYY-MM-DD; blank = today (Pakistan time)
PO_NOTES = ""                    # optional line printed under the table of every PO
PDF_DATE_FORMAT = "%d-%b"        # PDF name = <vendor>_<date>.pdf  ->  AAA WHOLESALE_13-Oct.pdf
                                 # other examples: "%d-%b-%Y" = 13-Oct-2026, "%d%m%y" = 131026, "%Y-%m-%d" = 2026-10-13
SUBFOLDER_BY_DATE = True         # True = files go into  <output_folder>\PO YYYY-MM-DD
SHOW_UPC_IN_PDF = False          # True = print box/master UPC under each description
SHOW_UPC_IN_EXCEL = True
CURRENCY = "$"
# =====================================================================================


def trim(value):
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip() if value is not None else ""


def to_number(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = trim(value).replace(",", "").replace(CURRENCY, "")
    try:
        return float(text) if text else None
    except ValueError:
        return None


def qty_text(value):
    return f"{value:,.0f}" if float(value).is_integer() else f"{value:,.2f}"


def price_text(value):
    text = f"{value:,.4f}".rstrip("0")
    if text.endswith("."):
        text += "00"
    elif len(text.split(".")[1]) < 2:
        text += "0" * (2 - len(text.split(".")[1]))
    return text


def money(value):
    return f"{CURRENCY}{value:,.2f}"


def safe_name(text):
    return re.sub(r"[\\/:*?\"<>|]+", " ", trim(text)).strip(" .")[:80] or "Vendor"


def read_marked_lines(path):
    """Returns (lines, skipped). A line is a dict; skipped rows carry a reason."""
    if PO_PRICE_BASIS not in ("last", "lowest"):
        raise ValueError('PO_PRICE_BASIS must be "last" or "lowest".')
    vendor_key = "last_vendor" if PO_PRICE_BASIS == "last" else "lowest_vendor"
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except PermissionError:
        raise PermissionError(f"Cannot open {path}. Close it in Excel and run again.")
    try:
        if ORDERING_SHEET not in workbook.sheetnames:
            raise ValueError(f"Sheet '{ORDERING_SHEET}' not found. Sheets: {', '.join(workbook.sheetnames)}")
        rows = workbook[ORDERING_SHEET].iter_rows(values_only=True)
        header = [trim(v) for v in next(rows, ())]
        lowered = {h.casefold(): i for i, h in enumerate(header) if h}

        def column(name, required=True):
            index = lowered.get(trim(name).casefold())
            if index is None and required:
                raise ValueError(f"Column '{name}' not found in '{ORDERING_SHEET}'. Headers: {', '.join(h for h in header if h)}")
            return index
        po_col = column(PO_COLUMN)
        c_desc, c_qty = column(COLUMNS["description"]), column(COLUMNS["qty"])
        c_vendor = column(COLUMNS[vendor_key])  # the price basis only decides which vendor column is used
        c_box, c_master = column(COLUMNS["box_upc"], False), column(COLUMNS["master_upc"], False)
        c_vover = column(COLUMNS["vendor_override"], False)
        mark = trim(PO_MARK).casefold()
        lines, skipped, marked, uncalculated = [], [], 0, 0
        for number, row in enumerate(rows, 2):
            get = lambda i: row[i] if i is not None and i < len(row) else None
            if trim(get(po_col)).casefold() != mark:
                continue
            marked += 1
            item = {"row": number, "description": trim(get(c_desc)), "qty": to_number(get(c_qty)),
                    "vendor": trim(get(c_vendor)),
                    "box": trim(get(c_box)), "master": trim(get(c_master))}
            if trim(get(c_vover)):
                item["vendor"] = trim(get(c_vover))   # typed vendor wins over Last/Lowest vendor
            if get(c_qty) is None:
                uncalculated += 1
            if item["qty"] is None or item["qty"] <= 0:
                item["reason"] = f"{COLUMNS['qty']} is blank or 0"
            elif not item["vendor"]:
                item["reason"] = f"no vendor ({COLUMNS[vendor_key]} is blank; type one in '{COLUMNS['vendor_override']}')"
            if item.get("reason"):
                skipped.append(item)
            else:
                lines.append(item)
        if marked and uncalculated == marked:
            raise ValueError("The marked rows have no calculated values. Open the ordering workbook in Excel, "
                             "save it (Ctrl+S) so the formulas keep their results, close it, and run again.")
        if not marked:
            raise ValueError(f"No row has '{PO_MARK}' in column '{PO_COLUMN}'.")
        return lines, skipped
    finally:
        workbook.close()


def group_by_vendor(lines):
    groups = OrderedDict()
    names = {}
    for item in sorted(lines, key=lambda x: (x["vendor"].casefold(), x["description"].casefold())):
        key = item["vendor"].casefold()
        names.setdefault(key, item["vendor"])
        groups.setdefault(key, []).append(item)
    return [(names[key], items) for key, items in groups.items()]


def pdf_file_name(vendor, po_date):
    return f"{safe_name(vendor)}_{po_date.strftime(PDF_DATE_FORMAT)}.pdf"


def write_pdf(path, po_no, po_date, vendor, items):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    normal = ParagraphStyle("n", parent=styles["Normal"], fontSize=9, leading=11)
    bold = ParagraphStyle("b", parent=normal, fontName="Helvetica-Bold")
    small = ParagraphStyle("s", parent=normal, fontSize=7.5, leading=9, textColor=colors.HexColor("#555555"))
    title = ParagraphStyle("t", parent=styles["Title"], fontSize=20, leading=24, alignment=2, textColor=colors.HexColor("#17365D"))
    right = ParagraphStyle("r", parent=normal, alignment=2)

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(15 * mm, 10 * mm, f"{po_no}  -  {vendor}")
        canvas.drawRightString(A4[0] - 15 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    document = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                                 topMargin=15 * mm, bottomMargin=18 * mm, title=f"{po_no} {vendor}", author=COMPANY["name"])
    company = [Paragraph(f"<b>{trim(COMPANY['name'])}</b>", ParagraphStyle("c", parent=normal, fontSize=12, leading=15))]
    company += [Paragraph(trim(line), normal) for line in COMPANY["lines"] if trim(line)]
    head = Table([[company, [Paragraph("PURCHASE ORDER", title)]]], colWidths=[95 * mm, 85 * mm])
    head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    info = Table([[Paragraph("<b>Vendor</b>", normal), Paragraph(vendor, bold), Paragraph("<b>PO No.</b>", normal), Paragraph(po_no, right)],
                  ["", "", Paragraph("<b>Date</b>", normal), Paragraph(po_date.strftime("%B %d, %Y"), right)]],
                 colWidths=[20 * mm, 85 * mm, 22 * mm, 53 * mm])
    info.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#17365D")),
                              ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F2F6FA"))]))

    data = [["#", "Description", "Qty"]]
    total_qty = 0.0
    for number, item in enumerate(items, 1):
        total_qty += item["qty"]
        text = Paragraph(item["description"], normal)
        if SHOW_UPC_IN_PDF and (item["box"] or item["master"]):
            codes = "  ".join(x for x in (f"Box: {item['box']}" if item["box"] else "", f"UPC: {item['master']}" if item["master"] else "") if x)
            text = [Paragraph(item["description"], normal), Paragraph(codes, small)]
        data.append([str(number), text, qty_text(item["qty"])])
    data.append(["", Paragraph("<b>TOTAL</b>", normal), Paragraph(f"<b>{qty_text(total_qty)}</b>", right)])
    table = Table(data, colWidths=[12 * mm, 138 * mm, 30 * mm], repeatRows=1)
    last = len(data) - 1
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (0, 0), (0, -1), "CENTER"), ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA5B1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, last - 1), [colors.white, colors.HexColor("#F7F9FB")]),
        ("BACKGROUND", (0, last), (-1, last), colors.HexColor("#E3ECF5")), ("LINEABOVE", (0, last), (-1, last), 1.2, colors.HexColor("#17365D")),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story = [head, Spacer(1, 6 * mm), info, Spacer(1, 6 * mm), table]
    if trim(PO_NOTES):
        story += [Spacer(1, 5 * mm), Paragraph(trim(PO_NOTES), normal)]
    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return total_qty


def write_merged_excel(path, po_date, orders, skipped):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "All PO lines"
    header_fill, header_font = PatternFill("solid", fgColor="17365D"), Font(bold=True, color="FFFFFF")
    edge = Side(style="thin", color="999999")
    border = Border(left=edge, right=edge, top=edge, bottom=edge)

    def header(ws, names, widths):
        ws.append(names)
        for index, (cell, width) in enumerate(zip(ws[1], widths), 1):
            cell.fill, cell.font, cell.border = header_fill, header_font, border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.row_dimensions[1].height = 24
        ws.sheet_view.showGridLines = False

    names = ["PO No.", "PO Date", "Vendor"] + (["box upc", "master upc"] if SHOW_UPC_IN_EXCEL else []) + ["Description", "Qty"]
    widths = [20, 13, 30] + ([18, 18] if SHOW_UPC_IN_EXCEL else []) + [60, 10]
    header(sheet, names, widths)
    grand_qty = 0.0
    desc_col = 4 + (2 if SHOW_UPC_IN_EXCEL else 0)
    for po_no, vendor, items, _ in orders:
        for item in items:
            row = [po_no, po_date, vendor] + ([item["box"], item["master"]] if SHOW_UPC_IN_EXCEL else []) + [item["description"], item["qty"]]
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                cell.border = border
            sheet.cell(sheet.max_row, 2).number_format = "yyyy-mm-dd"
            sheet.cell(sheet.max_row, desc_col).alignment = Alignment(wrap_text=True)
            sheet.cell(sheet.max_row, desc_col + 1).alignment = Alignment(horizontal="right")
            sheet.cell(sheet.max_row, desc_col + 1).number_format = "#,##0.##"
            grand_qty += item["qty"]
    sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"

    summary = workbook.create_sheet("PO summary")
    header(summary, ["PO No.", "PO Date", "Vendor", "Lines", "Total Qty", "PDF file"], [20, 13, 34, 8, 12, 60])
    for po_no, vendor, items, total_qty in orders:
        summary.append([po_no, po_date, vendor, len(items), total_qty, pdf_file_name(vendor, po_date)])
        for cell in summary[summary.max_row]:
            cell.border = border
        summary.cell(summary.max_row, 2).number_format = "yyyy-mm-dd"
        summary.cell(summary.max_row, 5).number_format = "#,##0.##"
    summary.append(["TOTAL", "", "", sum(len(o[2]) for o in orders), grand_qty, ""])
    for cell in summary[summary.max_row]:
        cell.font, cell.border, cell.fill = Font(bold=True), border, PatternFill("solid", fgColor="E3ECF5")
    summary.cell(summary.max_row, 5).number_format = "#,##0.##"

    if skipped:
        sk = workbook.create_sheet("Not ordered")
        header(sk, ["Row in ordering sheet", "Description", "Qty", "Vendor", "Reason"], [12, 60, 10, 30, 50])
        for item in skipped:
            sk.append([item["row"], item["description"], item["qty"], item["vendor"], item["reason"]])
            for cell in sk[sk.max_row]:
                cell.border = border
    workbook.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    for key, default in CONFIG.items():
        parser.add_argument("--" + key, default=default)
    args = parser.parse_args()
    source = Path(args.ordering_workbook).expanduser().resolve()
    if not source.is_file():
        print(f"[ERROR] Ordering workbook not found: {source}", file=sys.stderr)
        return 1
    try:
        po_date = date.fromisoformat(PO_DATE) if PO_DATE else datetime.now(timezone(timedelta(hours=5))).date()
        print(f"Reading {source}  (vendors: {PO_PRICE_BASIS.upper()}, or the typed PO vendor)", flush=True)
        lines, skipped = read_marked_lines(source)
        folder = Path(args.output_folder).expanduser().resolve()
        if SUBFOLDER_BY_DATE:
            folder = folder / f"PO {po_date.isoformat()}"
        folder.mkdir(parents=True, exist_ok=True)
        orders = []
        for number, (vendor, items) in enumerate(group_by_vendor(lines), 1):
            po_no = f"{PO_PREFIX}-{po_date:%Y%m%d}-{number:02d}"
            pdf = folder / pdf_file_name(vendor, po_date)
            total_qty = write_pdf(pdf, po_no, po_date, vendor, items)
            orders.append((po_no, vendor, items, total_qty))
            print(f"  {po_no}  {vendor}: {len(items)} lines, qty {qty_text(total_qty)}", flush=True)
        merged = folder / f"All_POs_{po_date.isoformat()}.xlsx"
        write_merged_excel(merged, po_date, orders, skipped)
        print(f"\n{len(orders)} PO(s), {len(lines)} lines, total qty {qty_text(sum(o[3] for o in orders))}")
        if skipped:
            print(f"WARNING: {len(skipped)} marked row(s) could not be ordered (see 'Not ordered' tab):")
            for item in skipped[:15]:
                print(f"  row {item['row']}: {item['description']} -> {item['reason']}")
        print(f"Folder: {folder}\nMerged Excel: {merged}")
        return 0
    except (ValueError, PermissionError, ImportError) as exc:
        print(f"[ERROR] {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
