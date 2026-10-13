"""
PDI Sales Sorter — filter and summarize store-level sales by department.

Reads one or more monthly PDI sales workbooks, filters rows whose Department
column contains a keyword (default: "novelty"), groups by description, and
writes a single summary workbook with one sheet per month.  Optionally writes
a detail workbook per month with the full merged and filtered rows.
"""

from pathlib import Path
from collections import OrderedDict
from functools import lru_cache
import json
import time
import sys
from decimal import Decimal, InvalidOperation
import re
import zipfile
from openpyxl import load_workbook, Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

HEADERS = ["Scan Code", "Description", "Qty Sold", "Retail", "Total Retail", "Department"]
LIMIT = 1_048_576


def text(value):
    return " ".join(str(value).split()) if value is not None else ""


def header_key(value):
    return re.sub(r"[^a-z0-9]", "", text(value).lower())


@lru_cache(maxsize=100000)
def scan_code(value):
    if value is None or text(value) == "":
        return "-"
    if isinstance(value, (int, float, Decimal)):
        value = format(Decimal(str(value)), "f")
    value = text(value)
    if re.fullmatch(r"\d+\.0+", value):
        value = value.split(".")[0]
    parts = re.split(r"[,;\s]+", value)
    return ", ".join(dict.fromkeys(
        (p.lstrip("0") or "0") if p.isdigit() else p for p in parts if p
    )) or "-"


def number(value, context, blank_zero=True):
    if value is None or text(value) == "":
        if blank_zero:
            return Decimal(0)
        return None
    s = text(value).replace(",", "").replace("$", "")
    if s.startswith("(") and s.endswith(")"):
        s = "-" + s[1:-1]
    try:
        result = Decimal(s)
        if not result.is_finite():
            raise InvalidOperation
        return result
    except InvalidOperation:
        raise ValueError(f"Invalid numeric value {value!r} at {context}")


def excel_number(value):
    return int(value) if value == value.to_integral_value() else float(value)


class OutputTable:
    """Streaming output, with transparent splitting at Excel's row limit."""
    def __init__(self, workbook, name, headers):
        self.workbook, self.name, self.headers = workbook, name, headers
        self.part = 0
        self.count = 0
        self.new_sheet()

    def new_sheet(self):
        self.part += 1
        suffix = "" if self.part == 1 else f" {self.part}"
        self.sheet = self.workbook.create_sheet(self.name[:31 - len(suffix)] + suffix)
        self.sheet.freeze_panes = "A2"
        widths = [36, 65, 18, 28, 20, 36, 18]
        for i in range(len(self.headers)):
            self.sheet.column_dimensions[chr(65 + i)].width = widths[i]
        cells = []
        for label in self.headers:
            c = WriteOnlyCell(self.sheet, label)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="24476A")
            cells.append(c)
        self.sheet.append(cells)
        self.rows = 1

    def append(self, values):
        if self.rows >= LIMIT:
            self.finish_sheet()
            self.new_sheet()
        cells = []
        for i, value in enumerate(values):
            if isinstance(value, str):
                c = WriteOnlyCell(self.sheet, value)
                c.data_type = "s"
                cells.append(c)
            else:
                cells.append(value)
        self.sheet.append(cells)
        self.rows += 1
        self.count += 1

    def finish_sheet(self):
        self.sheet.auto_filter.ref = f"A1:{chr(64 + len(self.headers))}{self.rows}"

    def finish(self):
        self.finish_sheet()


def locate_header(sheet):
    sheet.reset_dimensions()
    iterator = sheet.iter_rows(values_only=True)
    wanted = [header_key(h) for h in HEADERS]
    for rowno in range(1, 21):
        row = next(iterator, None)
        if row is None:
            break
        keys = [header_key(v) for v in row]
        if all(k in keys for k in wanted):
            return iterator, [keys.index(k) for k in wanted], rowno
    return None


def process_month(month, path, summary_book, output_dir, department_contains, save_detail):
    started = time.perf_counter()
    detail_path = output_dir / f"{month}_Merged_Filtered.xlsx"
    source = load_workbook(path, read_only=True, data_only=True)
    detail_book = Workbook(write_only=True) if save_detail else None
    merged = OutputTable(detail_book, "Merged", HEADERS + ["Month"]) if detail_book else None
    filtered = OutputTable(detail_book, "Filtered", HEADERS[:-1] + ["Month"]) if detail_book else None
    groups = {}
    counts = {"raw": 0, "filtered": 0, "blank_description": 0}
    source_qty, source_retail = Decimal(0), Decimal(0)
    included, skipped = [], []
    try:
        for sheet in source:
            if re.search(r"pivot|povit|\bpvt\b", sheet.title, re.I):
                skipped.append(sheet.title)
                continue
            located = locate_header(sheet)
            if located is None:
                skipped.append(sheet.title)
                continue
            included.append(sheet.title)
            iterator, indexes, header_row = located
            print(f"{month}: reading {sheet.title}", flush=True)
            for rowno, row in enumerate(iterator, header_row + 1):
                values = [row[i] if i < len(row) else None for i in indexes]
                if all(v is None or text(v) == "" for v in values):
                    continue
                if values[0] == "Scan Code" and values[1] == "Description":
                    continue
                code, description, qty, retail, total, department = values
                description, department = text(description), text(department)
                if description.casefold() in {"total", "grand total"} and not text(code):
                    continue
                context = f"{path.name}/{sheet.title}/row {rowno}"
                code = scan_code(code)
                qty = number(qty, context + " Qty Sold")
                total = number(total, context + " Total Retail")
                retail = number(retail, context + " Retail", blank_zero=False)
                cleaned = [code, description, excel_number(qty),
                           excel_number(retail) if retail is not None else None,
                           excel_number(total), department]
                counts["raw"] += 1
                if counts["raw"] % 50000 == 0:
                    elapsed = time.perf_counter() - started
                    print(f"  {month}: {counts['raw']:,} rows; {elapsed:.0f}s elapsed", flush=True)
                if merged:
                    merged.append(cleaned + [month])
                if department_contains.casefold() not in department.casefold():
                    continue
                counts["filtered"] += 1
                source_qty += qty
                source_retail += total
                if filtered:
                    filtered.append(cleaned[:-1] + [month])
                key = description.casefold()
                if not key:
                    key = (sheet.title, rowno)
                    counts["blank_description"] += 1
                group = groups.get(key)
                if group is None:
                    group = {"description": description, "codes": OrderedDict(),
                             "qty": Decimal(0), "total": Decimal(0), "prices": set()}
                    groups[key] = group
                for token in code.split(", "):
                    if token != "-":
                        group["codes"][token] = None
                group["qty"] += qty
                group["total"] += total
                if retail is not None:
                    group["prices"].add(retail)

        if not included:
            raise ValueError(f"{month}: no raw tab with the required six headers found")
        assert sum((g["qty"] for g in groups.values()), Decimal(0)) == source_qty
        assert sum((g["total"] for g in groups.values()), Decimal(0)) == source_retail
        summary = OutputTable(summary_book, month, HEADERS[:-1] + ["Month"])
        for group in sorted(groups.values(), key=lambda g: g["description"].casefold()):
            prices = sorted(group["prices"])
            price = (excel_number(prices[0]) if len(prices) == 1 else
                     ", ".join(format(p, "f") for p in prices) if prices else None)
            result_row = [", ".join(group["codes"]) or "-", group["description"],
                          excel_number(group["qty"]), price, excel_number(group["total"]), month]
            summary.append(result_row)
        summary.finish()
        if detail_book:
            merged.finish()
            filtered.finish()
            print(f"{month}: saving merged/filtered workbook...", flush=True)
            detail_book.save(detail_path)
        report = (f"{month}: raw={counts['raw']:,}; filtered={counts['filtered']:,}; "
                  f"summary={len(groups):,}; Qty Sold={source_qty}; Total Retail={source_retail}; "
                  f"blank descriptions={counts['blank_description']}; "
                  f"included={included}; skipped={skipped}")
        print(report, flush=True)
        return report
    finally:
        source.close()


def run(input_files, output_dir, department_contains="novelty", save_detail=False):
    """
    Parameters
    ----------
    input_files : dict[str, str]
        {month_label: file_path} — one or more monthly PDI sales files.
    output_dir : str
        Directory for output files.
    department_contains : str
        Case-insensitive substring filter on the Department column.
    save_detail : bool
        If True, also write per-month Merged+Filtered detail workbooks.

    Returns
    -------
    dict with keys:
        "summary" : str — path to the summary workbook
        "details" : list[str] — paths to detail workbooks (empty if save_detail is False)
        "reports" : list[str] — per-month report strings
    """
    # Validate input files
    errors = []
    for month, filename in input_files.items():
        path = Path(filename)
        try:
            if not path.is_file():
                raise FileNotFoundError(filename)
            with zipfile.ZipFile(path) as archive:
                if "xl/workbook.xml" not in archive.namelist():
                    raise ValueError("not an Excel XLSX workbook")
        except Exception as exc:
            errors.append(f"{month}: {exc}")
    if errors:
        raise ValueError("Input file errors:\n" + "\n".join(errors))

    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    summary_book = Workbook(write_only=True)
    reports = []
    for month, filename in input_files.items():
        reports.append(process_month(month, Path(filename), summary_book, out,
                                     department_contains, save_detail))

    summary_path = out / "Stores_Sales_Summary.xlsx"
    summary_book.save(summary_path)
    print(f"\nFinished: {summary_path}", flush=True)

    detail_files = []
    if save_detail:
        for month in input_files:
            dp = out / f"{month}_Merged_Filtered.xlsx"
            if dp.is_file():
                detail_files.append(str(dp))

    return {
        "summary": str(summary_path),
        "details": detail_files,
        "reports": reports,
    }
